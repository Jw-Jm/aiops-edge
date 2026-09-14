from pathlib import Path
from datetime import date

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor, Twips


ROOT = Path("/Users/mssc/Documents/Code/agent/aiops")
IMG = ROOT / "docs/design/ui-renders"
OUT = ROOT / "docs/AIOps平台最终设计与实施规范.docx"

BLUE = "245B89"
DARK_BLUE = "24364B"
PALE_BLUE = "EDF4FA"
PALE_GRAY = "F7F9FC"
MID_GRAY = "667085"
LIGHT_BORDER = "D9D9D9"
BLACK = "111827"
RED = "B42318"
GREEN = "027A48"
AMBER = "B54708"
BODY_FONT = "Arial Unicode MS"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color=LIGHT_BORDER, size="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=80, start=100, bottom=80, end=100):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_width(cell, inches):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(inches * 1440)))
    tc_w.set(qn("w:type"), "dxa")


def apply_exact_table_geometry(table, widths):
    """Keep tblW, tblGrid, column widths and every tcW synchronized."""
    if widths:
        column_widths = [int(round(value * 1440)) for value in widths]
    else:
        total_width = int(round(6.6 * 1440))
        column_widths = [total_width // len(table.columns)] * len(table.columns)
        column_widths[-1] += total_width - sum(column_widths)
    table_width = sum(column_widths)

    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr

    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:type"), "dxa")
    tbl_w.set(qn("w:w"), str(table_width))

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:type"), "dxa")
    tbl_ind.set(qn("w:w"), "100")

    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in column_widths:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(width))
        grid.append(grid_col)

    for column_index, width in enumerate(column_widths):
        table.columns[column_index].width = Twips(width)
        for row in table.rows:
            cell = row.cells[column_index]
            cell.width = Twips(width)
            tc_w = cell._tc.get_or_add_tcPr().find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                cell._tc.get_or_add_tcPr().append(tc_w)
            tc_w.set(qn("w:type"), "dxa")
            tc_w.set(qn("w:w"), str(width))


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_begin, instr, fld_sep, fld_end])


def add_text(paragraph, text, bold=False, color=BLACK, size=None, italic=False, font=BODY_FONT):
    run = paragraph.add_run(text)
    run.bold = bold
    run.italic = italic
    run.font.name = font
    run._element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
    run.font.color.rgb = RGBColor.from_string(color)
    if size:
        run.font.size = Pt(size)
    return run


def add_body(doc, text, bold_lead=None, style=None, keep=False):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.18
    p.paragraph_format.keep_together = keep
    if bold_lead and text.startswith(bold_lead):
        add_text(p, bold_lead, bold=True)
        add_text(p, text[len(bold_lead):])
    else:
        add_text(p, text)
    return p


def add_bullets(doc, items, level=0):
    for item in items:
        p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
        p.paragraph_format.left_indent = Inches(0.23 + level * 0.2)
        p.paragraph_format.first_line_indent = Inches(-0.15)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.12
        add_text(p, item)


def add_numbered(doc, items):
    # Render numbering explicitly so every independent list restarts at 1.
    # Word's built-in List Number style otherwise shares numbering state across
    # distant sections when the document is regenerated by python-docx.
    for index, item in enumerate(items, start=1):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.28)
        p.paragraph_format.first_line_indent = Inches(-0.18)
        p.paragraph_format.space_after = Pt(1.5)
        p.paragraph_format.line_spacing = 1.0
        add_text(p, f"{index}. ", bold=True)
        add_text(p, item)


def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    p.paragraph_format.keep_with_next = True
    return p


def add_table(doc, headers, rows, widths=None, font_size=8.8, header_fill=DARK_BLUE):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.style = "Table Grid"
    header = table.rows[0]
    set_repeat_table_header(header)
    prevent_row_split(header)
    for i, value in enumerate(headers):
        cell = header.cells[i]
        set_cell_shading(cell, header_fill)
        set_cell_border(cell)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        if widths:
            set_width(cell, widths[i])
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        add_text(p, str(value), bold=True, color="FFFFFF", size=font_size)
    for ridx, row_values in enumerate(rows):
        row = table.add_row()
        prevent_row_split(row)
        if ridx % 2 == 1:
            for cell in row.cells:
                set_cell_shading(cell, PALE_GRAY)
        for i, value in enumerate(row_values):
            cell = row.cells[i]
            set_cell_border(cell)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            if widths:
                set_width(cell, widths[i])
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            add_text(p, str(value), size=font_size)
    apply_exact_table_geometry(table, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_figure(doc, filename, caption, alt, width=6.8, page_break=False):
    if page_break:
        doc.add_page_break()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    run = p.add_run()
    shape = run.add_picture(str(IMG / filename), width=Inches(width))
    shape._inline.docPr.set("descr", alt)
    cp = doc.add_paragraph()
    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cp.paragraph_format.keep_with_next = True
    cp.paragraph_format.space_after = Pt(6)
    add_text(cp, caption, italic=True, color=MID_GRAY, size=9)


def add_page_note(doc, title, items):
    p = doc.add_paragraph()
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_after = Pt(2)
    add_text(p, title, bold=True, color=BLUE, size=10.5)
    add_bullets(doc, items)


def configure_document(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.68)
    section.bottom_margin = Inches(0.68)
    section.left_margin = Inches(0.72)
    section.right_margin = Inches(0.72)
    section.header_distance = Inches(0.25)
    section.footer_distance = Inches(0.28)
    section.different_first_page_header_footer = True

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = BODY_FONT
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(BLACK)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.18

    title = styles["Title"]
    title.font.name = BODY_FONT
    title._element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
    title.font.size = Pt(27)
    title.font.bold = True
    title.font.color.rgb = RGBColor.from_string(BLACK)
    title.paragraph_format.space_after = Pt(10)
    title_ppr = title._element.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)

    for level, size in ((1, 18), (2, 14), (3, 11.5), (4, 10.5)):
        style = styles[f"Heading {level}"]
        style.font.name = BODY_FONT
        style._element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(BLACK)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(12 if level == 1 else 8)
        style.paragraph_format.space_after = Pt(5)
        if level == 1:
            style.paragraph_format.page_break_before = True

    for style_name in ("List Bullet", "List Bullet 2", "List Number"):
        style = styles[style_name]
        style.font.name = BODY_FONT
        style._element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
        style.font.size = Pt(10.3)
        style.font.color.rgb = RGBColor.from_string(BLACK)

    if "Subtitle AIOps" not in styles:
        subtitle = styles.add_style("Subtitle AIOps", WD_STYLE_TYPE.PARAGRAPH)
    else:
        subtitle = styles["Subtitle AIOps"]
    subtitle.font.name = BODY_FONT
    subtitle._element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
    subtitle.font.size = Pt(14)
    subtitle.font.color.rgb = RGBColor.from_string(MID_GRAY)
    subtitle.paragraph_format.space_after = Pt(18)

    # Minimal running header and footer.
    hp = section.header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hp.paragraph_format.space_after = Pt(0)
    add_text(hp, "AIOps 平台最终设计与实施规范  ·  V1.4", color=MID_GRAY, size=8)
    fp = section.footer.paragraphs[0]
    add_page_number(fp)
    for run in fp.runs:
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor.from_string(MID_GRAY)

    settings = doc.settings.element
    update_fields = OxmlElement("w:updateFields")
    update_fields.set(qn("w:val"), "true")
    settings.append(update_fields)


def add_cover(doc):
    for _ in range(3):
        doc.add_paragraph()
    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_text(p, "AIOps 平台最终设计与实施规范", bold=True, size=27)
    p = doc.add_paragraph(style="Subtitle AIOps")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_text(p, "云平台智能运维产品 UI 数据 图谱 智能能力和验收基线", color=MID_GRAY, size=14)
    doc.add_paragraph()
    meta = [
        ("文档版本", "V1.4（八页最终信息架构版）"),
        ("目标系统", "http://localhost:30253/"),
        ("设计日期", "2026-09-13"),
        ("适用对象", "产品、前端、后端、AI、测试与验收智能体"),
        ("规范强度", "本文的“必须 / 不得 / 应”均为可验收要求"),
    ]
    table = add_table(doc, ["项目", "内容"], meta, widths=[1.35, 5.25], font_size=10, header_fill=BLUE)
    table.rows[0].cells[0].paragraphs[0].clear()
    table.rows[0].cells[1].paragraphs[0].clear()
    set_cell_shading(table.rows[0].cells[0], "FFFFFF")
    set_cell_shading(table.rows[0].cells[1], "FFFFFF")
    for cell in table.rows[0].cells:
        set_cell_border(cell, "FFFFFF", "0")
    doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(10)
    add_text(p, "核心产品决策", bold=True, color=BLUE, size=12)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.left_indent = Inches(0.45)
    p.paragraph_format.right_indent = Inches(0.45)
    p.paragraph_format.space_after = Pt(8)
    add_text(p, "统一账户、八个一级页面。AI 智能运维置于第一入口；其后依次为总览、集群、全链路监控、知识图谱、知识库、报告和设置。告警、资源、调查过程与处置记录按上下文嵌入上述页面，不再设独立入口。", bold=True, size=12)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.left_indent = Inches(0.6)
    p.paragraph_format.right_indent = Inches(0.6)
    add_text(p, "UI 图中的数值仅用于表达布局和异常态，不构成验收数据。验收必须连接真实环境并形成可复核的数据来源链。", color=MID_GRAY, size=9.5)


def build_doc():
    doc = Document()
    configure_document(doc)
    props = doc.core_properties
    props.title = "AIOps 平台最终设计与实施规范"
    props.subject = "云平台智能运维产品最终 UI 与实施验收基线"
    props.author = "AIOps 项目组"
    props.keywords = "AIOps, UI, 云平台, 知识图谱, LLM, MCP, 工作流, 验收"
    add_cover(doc)

    add_heading(doc, "文档导航", 1)
    add_body(doc, "本文是实现与验收的唯一产品基线。智能体应按章节编号逐项执行，并使用 Word 导航窗格定位标题。任何偏离必须记录为设计变更，不得以“现有页面如此”作为保留理由。")
    nav = [
        "1  文档用途 范围和术语",
        "2  不可变设计决策",
        "3  产品能力边界与完整性模型",
        "4  信息架构 导航和路由",
        "5  全局交互和视觉规范",
        "6  页面级最终设计",
        "7  云平台运维信息完整性",
        "8  后端能力的前端归属与覆盖规则",
        "9  知识图谱自动更新 运维和展示",
        "10 设置中的 AI 能力配置",
        "11 AI 智能故障定位能力和合理性分析",
        "12 真实数据 状态语义和可追溯性",
        "13 统一账户下的安全与危险操作",
        "14 响应式 可访问性 性能和可观测性",
        "15 分阶段实施任务和代码落点",
        "16 完整验收门禁和交付物",
        "17 其他智能体的执行提示词",
        "附录 A 路由和旧入口迁移表",
        "附录 B 页面级验收检查表",
        "附录 C 需求追踪矩阵和最终完成定义",
    ]
    nav_rows = list(zip(nav[:10], nav[10:]))
    add_table(doc, ["章节 1—10", "章节 11—附录"], nav_rows, widths=[3.3, 3.3], font_size=8.6)
    add_body(doc, "阅读顺序：产品/设计智能体先读 1—7；前后端智能体必须读 1—15；测试/验收智能体必须完整阅读全文。", bold_lead="阅读顺序：")

    add_heading(doc, "1  文档用途 范围和术语", 1)
    add_heading(doc, "1.1 用途", 2)
    add_body(doc, "本规范把产品目标、信息架构、页面结构、数据契约、后端覆盖、知识图谱、AI 控制面、智能故障定位、安全边界、实现顺序和验收证据固化为同一份可执行文档。它不是概念稿，也不是已有页面的评审记录，而是从零设计后的目标状态。")
    add_heading(doc, "1.2 范围", 2)
    add_bullets(doc, [
        "覆盖目标站点所有可见页面、交互、空态、异常态、加载态、权限/范围态和真实数据展示。",
        "覆盖所有公共后端能力的前端归属；内部编排与内部查询端点仅需可观测，不要求一端点一页面。",
        "覆盖云平台五类运维域：计算、网络、存储、Kubernetes/容器、应用与中间件，以及它们的因果关系。",
        "覆盖 LLM、Agent、Skill、Workflow、MCP、知识库、知识图谱、策略、安全、评测和运行可观测性。",
        "覆盖智能故障定位从触发、证据、假设、验证、结论到建议、执行与复核的完整闭环。",
    ])
    add_heading(doc, "1.3 术语与规范用语", 2)
    add_table(doc, ["术语", "本规范定义"], [
        ("统一账户", "同一已认证账户看到同一套页面与能力；不按值班人员/管理员拆导航。"),
        ("范围", "平台、租户、集群、命名空间、服务、资源、时间窗等数据与操作边界。"),
        ("运维任务", "在 AI 智能运维页内承载的一次完整故障定位或运维请求；不是独立一级页面。"),
        ("调查运行", "运维任务内部固定范围和数据版本后，收集证据、形成假设、验证因果并输出结论的运行实例。"),
        ("证据", "带来源、时间、范围、查询条件、摘要和可打开定位链接的不可含糊事实。"),
        ("公共后端能力", "面向前端或外部调用者的稳定 API；必须在页面、上下文动作或能力中心有归属。"),
        ("内部能力", "仅由服务间编排使用的端点；可不直接展示，但其运行健康必须可观测。"),
        ("真实数据", "从目标环境真实数据源读取、可回查、非前端 mock/fixture/静态构造的数据。"),
    ], widths=[1.35, 5.25])

    add_heading(doc, "2  不可变设计决策", 1)
    add_table(doc, ["ID", "决策", "验收判定"], [
        ("D-01", "统一账户；不拆值班人员与管理员产品。", "任意登录后看到同一套八项主导航和设置；不得出现 adminOnly 导航隐藏。"),
        ("D-02", "八页信息架构固定。", "主导航严格为 AI 智能运维、总览、集群、全链路监控、知识图谱、知识库、报告、设置；顺序固定。"),
        ("D-03", "AI 智能运维是第一入口。", "默认首页进入 AI 智能运维；它承载提问、任务进度、证据、假设、结论、动作草稿、执行记录和恢复验证。"),
        ("D-04", "事实、推断和建议分层。", "事实显示来源；推断显示置信度与支持/反证；建议显示风险和预期效果。三者不得混排成一句无来源结论。"),
        ("D-05", "范围是第一等对象。", "顶栏永久显示当前平台/集群/资源/时间；调查启动后冻结范围并显示快照时间与图谱 generation。"),
        ("D-06", "真实数据验收。", "验收态禁止 mock、fixture、随机数、硬编码“健康”；所有关键数值可回查源接口或源系统。"),
        ("D-07", "高风险动作受控但不按角色拆分。", "写操作始终显示预检、影响范围、确认、审计和回滚；AI 不得自动批准或执行。"),
        ("D-08", "一能力一权威入口。", "旧路由可重定向，但不得形成同能力多套交互或重定向到无效 tab。"),
        ("D-09", "异常诚实可见。", "未知、陈旧、部分、未接入、失败均为显式状态；不得以 0 或绿色掩盖。"),
        ("D-10", "公共后端能力必须有前端归属。", "覆盖清单中每个公共 API 均映射到页面/动作/设置；内部 API 标注 API-only。"),
    ], widths=[0.55, 3.0, 3.05], font_size=8.1)
    add_heading(doc, "2.1 禁止项", 2)
    add_bullets(doc, [
        "不得恢复“值班人员端 / 管理员端”或按身份隐藏完整导航。",
        "不得新增问题、资源、调查、处置等一级或二级导航；这些能力必须嵌入八个既定页面。",
        "不得使用前端静态 JSON、演示数字或随机数据证明真实环境通过。",
        "不得把后端返回 200 视为功能完成；必须验证业务结果、前端状态、审计和回查。",
        "不得提供可点击但无效果的按钮、重定向到不存在的视图、无法回退的深链。",
        "不得让 LLM 直接持有基础设施凭据、绕过后端策略或执行未经确认的写操作。",
    ])

    add_heading(doc, "3  产品能力边界与完整性模型", 1)
    add_heading(doc, "3.1 产品闭环", 2)
    add_body(doc, "产品不是监控看板集合，而是围绕“发现—理解—定位—决策—执行—验证—沉淀”的统一运维闭环。AI 智能运维是闭环工作台；其他页面提供真实态势、证据、关系、知识和报告，并把所选上下文带回该工作台。")
    add_table(doc, ["阶段", "核心问题", "权威页面", "必须产物"], [
        ("发现", "哪些集群存在告警或关键指标异常？", "总览 / 集群", "集群告警、关键指标、范围、时间、可信状态"),
        ("理解", "异常位于哪条服务链和哪些上下游？", "集群 / 全链路监控 / 知识图谱", "资源身份、黄金信号、链路、关系、变更、影响面"),
        ("定位", "最可能的因果链是什么？", "AI 智能运维", "冻结范围、证据、候选假设、反证、结论等级"),
        ("决策", "采取什么动作、风险多大？", "AI 智能运维", "动作建议、预检、风险、显式确认信息"),
        ("执行", "动作实际发生了什么？", "AI 智能运维", "确定性执行记录、标准输出、审计、回滚信息"),
        ("验证", "问题是否恢复、有无副作用？", "AI 智能运维 / 报告", "前后指标、SLO、受影响范围复核"),
        ("沉淀", "如何让下一次更快更准？", "知识库 / 报告 / 设置", "运行手册、报告、评测集、规则和工作流版本"),
    ], widths=[0.65, 1.85, 1.15, 2.95])
    add_heading(doc, "3.2 云平台能力域", 2)
    add_table(doc, ["能力域", "必须覆盖", "不允许的降级"], [
        ("资产与拓扑", "平台/集群/命名空间/节点/VM/网络/存储/应用身份、关系、生命周期与负责人", "仅列 Kubernetes 容器；把网络、存储、硬件藏进原始 JSON"),
        ("可观测信号", "指标、日志、追踪、事件、告警、SLO、变更、容量", "只展示图表，无原查询、时间、单位、缺口和关联"),
        ("可靠性", "健康、可用性、容量风险、备份/恢复状态、证书与关键依赖", "以“无告警”等价于健康"),
        ("连续性与生命周期", "备份、恢复、灾备、升级、证书、补丁、配置漂移和变更窗口", "只显示配置存在，不验证最近成功和恢复结果"),
        ("安全与合规", "漏洞、暴露面、身份范围、危险动作、凭据轮换、策略偏差和审计", "用前端隐藏代替服务端控制；暴露 secret"),
        ("成本与效率", "成本分摊、预算、利用率、闲置、规格优化建议及其数据置信度", "把估算显示为账单事实；无 scope 或价格版本"),
        ("事件闭环", "问题、调查、处置、审计、复盘、知识沉淀", "问题和动作互不关联；执行后无法验证"),
        ("智能能力", "模型、Agent、Skill、Workflow、MCP、RAG、图谱、策略、评测、运行跟踪", "只有聊天框和 LLM 配置"),
        ("平台自身", "采集、查询、对象存储、图谱、AI 编排、外部依赖、缓存和队列健康", "AIOps 自身故障不可见"),
        ("治理", "账户、范围、凭据引用、数据外发、危险操作、安全审计", "在浏览器保存密钥；依赖前端做权限判断"),
    ], widths=[1.1, 3.4, 2.1])
    add_heading(doc, "3.3 发布优先级", 2)
    add_body(doc, "P0 是本次最终设计的发布门槛；P1 必须进入同一信息架构并显示真实接入状态；P2 可分期实现，但不得伪装为已支持。")
    add_table(doc, ["级别", "范围", "门禁"], [
        ("P0", "统一壳层、AI 智能运维、总览、集群、全链路监控、知识图谱、知识库、报告、设置、真实数据语义和安全闭环", "八个一级页面或 AI 定位—处置—验证闭环任一失败均阻断发布"),
        ("P1", "SLO、容量、静默、LLM/Agent/Workflow/MCP/RAG/评测、平台健康、备份/恢复状态、证书、配置漂移、补丁/漏洞清单", "公共后端已有能力必须全部归入八页之一；缺少数据源时显示真实未接入状态和影响"),
        ("P2", "灾备编排、自动补丁、FinOps 优化、外部 ITSM/On-call 联动", "进入能力目录和路线图；若纳入合同或已有公共后端则提升为 P0/P1"),
    ], widths=[0.65, 4.35, 1.6])

    add_heading(doc, "4  信息架构 导航和路由", 1)
    add_heading(doc, "4.1 八个一级页面", 2)
    add_body(doc, "主导航严格使用下列八项，AI 智能运维固定置顶并作为默认首页。告警在总览/集群中呈现；资源在集群、全链路和图谱中呈现；调查与处置作为 AI 任务内部阶段呈现。所有入口对统一账户可见，后端仍根据令牌、租户和资源范围校验请求。")
    add_table(doc, ["层级", "入口", "唯一职责", "规范路由"], [
        ("主入口", "AI 智能运维", "自然语言入口与定位、建议、受控执行、验证全闭环", "/ai-operations"),
        ("观察", "总览", "汇总全部纳管集群的告警、关键指标和数据质量", "/overview"),
        ("观察", "集群", "展示指定集群的计算、Kubernetes、网络、存储和活动告警", "/clusters/:clusterUid/overview"),
        ("观察", "全链路监控", "一页关联云平台控制面、计算、网络、存储和 Kubernetes 路径", "/observability"),
        ("观察", "知识图谱", "中心对象、有限子图、依赖路径、影响分析和 generation", "/knowledge-graph"),
        ("知识", "知识库", "知识源、运行手册、索引、检索验证和任务沉淀", "/knowledge"),
        ("输出", "报告", "巡检报告与 AI 智能运维任务完成报告", "/reports"),
        ("治理", "设置", "接入、图谱、LLM、Workflow、MCP、RAG、策略、安全和平台健康", "/settings"),
    ], widths=[0.85, 0.85, 3.6, 1.25])
    add_heading(doc, "4.2 顶栏与全局状态", 2)
    add_bullets(doc, [
        "左侧：页面标题和简短任务说明；不重复面包屑噪声。",
        "中部：全局范围选择器，依次为平台/租户、集群、命名空间或资源、时间窗。",
        "右侧：页面数据截止时间与时区、聚合质量状态、刷新、通知、帮助、统一账户菜单；不得出现角色切换器。",
        "页面级新鲜度不得只写“几分钟前”：必须说明该时间是数据截止、最近成功还是请求完成；不同来源不同步时显示最差质量并允许展开来源明细。",
        "范围变化必须更新 URL、取消旧请求、清除不兼容选择，并在 300 ms 内给出加载反馈。",
        "AI 任务运行时显示“冻结范围”；用户改变全局范围不得悄然改变正在运行的任务。",
    ])
    add_heading(doc, "4.3 深链与返回规则", 2)
    add_table(doc, ["规则", "要求"], [
        ("URL 可复现", "scope、time、tab、filter、selection 使用稳定查询参数；刷新后恢复。"),
        ("对象链接", "资源、证据、任务和动作记录均使用 canonical UID；显示名变化不破坏链接。"),
        ("来源返回", "从任一观察页进入 AI 智能运维时保留 sourcePage、scope、time、selection 与返回入口。"),
        ("无效参数", "显示明确的参数无效说明并回到安全默认视图；不得静默跳到无关页面。"),
        ("旧路由", "迁移期 301/replace 到规范路由且保留可转换参数；无法转换时显示迁移说明。"),
    ], widths=[1.2, 5.4])
    add_heading(doc, "4.4 设置页分区", 2)
    add_table(doc, ["子路由", "页面", "必须保留的状态"], [
        ("/settings?section=integrations", "接入与目录", "connector、scope、status、runId"),
        ("/settings?section=graph", "图谱运维", "source、generation、schema、runId"),
        ("/settings?section=llm", "LLM 与路由", "model/route ID、version、probeRunId"),
        ("/settings?section=workflow", "Agent 与 Workflow", "workflow/agent ID、version、runId"),
        ("/settings?section=mcp", "MCP 与工具", "server/tool ID、risk、version、runId"),
        ("/settings?section=rag", "知识与 RAG", "index/source ID、version、runId"),
        ("/settings?section=security", "策略、安全与范围", "policy version、scope、auditId"),
        ("/settings?section=health", "平台自身健康", "capability、component、time、probeRunId"),
    ], widths=[2.15, 1.45, 3.0], font_size=8.2)

    add_heading(doc, "5  全局交互和视觉规范", 1)
    add_heading(doc, "5.1 视觉方向", 2)
    add_body(doc, "目标是高信息密度但低认知负担的运维工作台：浅色中性背景、稳定左侧导航、内容优先、少量蓝色强调、状态色只表达状态。不得使用大面积渐变、装饰性 3D、无意义大卡片或把关键数据埋进抽屉。")
    add_table(doc, ["项目", "规范"], [
        ("画布", "桌面基准 1440×900；侧栏 184 px；顶栏 56–60 px；内容左右内边距 24 px。"),
        ("字体", "中文优先 Microsoft YaHei/系统无衬线；正文 14 px，表格 13–14 px，标题 20/16 px。"),
        ("颜色", "主色 #245B89；正文 #111827；次要 #667085；背景 #F7F9FC；边框 #D9D9D9。"),
        ("状态", "严重/失败红；风险/陈旧琥珀；健康绿；未知/未接入灰；不得仅靠颜色区分。"),
        ("圆角", "卡片 8 px，按钮 6 px，标签 999 px；同层级保持一致。"),
        ("阴影", "默认无阴影；浮层和拖拽对象使用低强度阴影。"),
        ("数据密度", "表格行 40–44 px；关键 KPI 不超过首屏 6 个；高基数信息进入表格/详情。"),
        ("图表", "轴、单位、时间窗、采样/聚合、缺失段、阈值必须可见；不使用无刻度“好看曲线”。"),
    ], widths=[1.15, 5.45])
    add_heading(doc, "5.2 组件与交互语义", 2)
    add_table(doc, ["组件", "必须包含", "禁止"], [
        ("KPI", "值、单位、口径、时间窗、变化方向、状态", "只有大数字；用 0 代替未知"),
        ("状态标签", "文本+图标/形状+颜色；可解释 tooltip", "只用红绿点"),
        ("数据表", "排序、过滤、列设置、固定主列、分页/虚拟化、空态", "横向溢出截断关键列"),
        ("详情抽屉", "稳定 URL、可复制链接、固定主动作、关闭后恢复列表位置", "承载完整复杂工作流"),
        ("确认对话框", "动作、对象、范围、影响、风险、前置条件、回滚", "“确定吗？”一行文字"),
        ("加载态", "结构骨架或明确进度、请求上下文、可取消长任务", "页面空白或无限转圈"),
        ("错误态", "可读原因、request/run ID、重试、文档/诊断入口", "展示堆栈或笼统“失败”"),
        ("空态", "区分真空、过滤为空、未接入、无权限、查询失败", "统一显示“暂无数据”"),
    ], widths=[1.05, 3.45, 2.1])
    add_heading(doc, "5.3 AI 智能运维入口规范", 2)
    add_bullets(doc, [
        "AI 智能运维是默认首页；其他页面固定提供“交给 AI 分析”上下文动作，动作进入同一 AI 页面而非打开第二套助手。",
        "上下文头必须显示：来源页面、scope、时间窗、选中对象、task/run ID、图谱 generation。",
        "每条 AI 结论使用“结论等级、置信度、支持证据、反证/缺口、下一步”结构。",
        "证据引用必须可点击并回到具体图表区间、日志查询、追踪 span、事件或资源关系。",
        "任何执行建议先生成动作草稿；预检、影响范围、显式确认、执行、回滚与恢复验证均在当前 AI 任务中完成。",
    ])

    add_heading(doc, "5.4 信息有效性与操作效率", 2)
    add_body(doc, "页面中的每一项信息必须帮助用户判断状态、定位原因、决定下一步或验证结果。不能回答上述任一问题的计数、卡片、标签、技术配置和装饰文本默认删除或移入详情。")
    add_table(doc, ["对象", "有效信息合同", "必须删除或降级"], [
        ("导航徽标", "仅显示待处理、失败、待确认或未读，并绑定明确查询、scope 与时间窗", "无口径总数、健康数量、与页面 KPI 重复的数字"),
        ("KPI/摘要", "metric_id、来源、分母、scope、时间窗、截止时间、质量、趋势和下钻", "只有大数字、100% 虚荣指标、用 0 代替 unknown/not connected"),
        ("时间", "审计/报告/计划任务使用完整日期、时间和时区；相对时间只作辅文", "只写 13:50、02:00 或“刚刚”却无法判断所属日期"),
        ("列表/卡片", "异常、部分、陈旧、失败优先；健康项折叠为可展开汇总", "逐项平铺大量相同“成功”卡片、无动作的绿色占位"),
        ("图谱", "同时显示返回总量、当前渲染量、过滤条件、边数、generation 和展开动作", "声明 16 个节点却只画 5 个且不解释"),
        ("AI 证据", "evidence_id、类型、来源时间、新鲜度/质量、范围和可打开定位", "只有自然语言陈述、没有原始事实链接或缺口动作"),
        ("安全范围", "只返回当前身份可见范围；越界只给拒绝、诊断 ID 和审计入口", "在“不可见”行中泄露其他租户、集群或对象名称"),
        ("不可用状态", "Not connected/Failed、失败阶段、最后成功、request/run ID、重试或接入动作", "沿用历史值冒充当前值、空白页面、笼统“暂无数据”"),
    ], widths=[1.05, 3.75, 1.8], font_size=8.0)
    add_bullets(doc, [
        "首屏采用异常优先和渐进披露：先显示需决策的异常、影响与主动作；健康明细、原始配置和诊断 JSON 下沉。",
        "按钮文案使用动词+对象，例如“发起全量对账”“运行只读探测”；不得使用含糊的“处理”“操作”或让普通查看按钮暗示直接发布。",
        "同一事实只保留一个权威位置，其他页面以链接和上下文摘要引用；任何重复值必须来自同一字段合同。环境不可达时只记录本次探测事实和最后成功记录，不得把历史截图或上次验收数值显示为当前状态。",
    ])

    add_heading(doc, "6  页面级最终设计", 1)
    add_body(doc, "本章 UI 图是布局、层级和状态表达的强约束；图中值只用于说明版式，不是当前环境观察或验收事实。生产实现必须由 API 返回运行值，并显示来源、口径、范围、时间与质量；允许响应式重排，不允许改变信息优先级、页面职责或安全语义。")

    page_specs = [
        ("6.1 AI 智能运维", "01-ai-operations.png", "图 1  AI 智能运维：故障定位、证据、建议与受控执行", "AI 智能运维主页面，包含范围输入、任务阶段、证据、候选原因、结论、建议动作和执行生命周期。", [
            "该页是导航第一项和默认首页；既可自由提问，也可从其他页面携带 scope、time、selected object/path/relation 进入。",
            "三栏结构固定：左侧任务与阶段，中部对话/分析及证据，右侧候选原因、结论等级、动作草稿和执行状态；窄屏按相同优先级纵向重排。",
            "运行开始后冻结 scope、时间窗、资源集合、selectedPathId/selectedRelationId、graph generation 及知识版本；变更上下文必须创建新任务版本。",
            "每条证据显示 evidence_id、类型、来源时间、范围、质量和可打开定位；事实、推断、反证/缺口和建议必须视觉分层。",
            "结论等级固定 Unknown/Candidate/Supported/Confirmed；证据不足、冲突或来源 Partial/Stale 时不得输出 Confirmed。",
            "动作状态在本页完成 Draft → Preflight → Awaiting confirmation → Running → Succeeded/Failed → Verified/Rolled back；AI 不得自动确认或持有基础设施凭据。",
            "任务完成可一键生成 AI 运维报告，报告绑定 task/run、全部证据、动作、前后验证和版本组合。",
        ]),
        ("6.2 总览", "02-platform-overview.png", "图 2  总览：全部纳管集群告警与跨集群关键指标", "AIOps 总览页面，汇总当前账户可见的全部纳管集群、告警态势、集群级计算指标和平台数据质量。", [
            "首屏回答：纳管了哪些集群、哪些集群有告警、异常集中在哪个资源域、数据是否可信。",
            "CPU 和内存均为集群口径：CPU 使用核数/可分配核数，内存已用/可分配容量，同时显示 P95 节点利用率或热点节点数，不能用某一节点代替集群。",
            "使用图表展示集群告警分布、跨集群 CPU/内存/节点就绪比较、容量与数据新鲜度；轴、单位、聚合、阈值、时间窗与缺失段必须可见。",
            "集群清单必须覆盖所有纳管集群；默认按严重告警、警告、数据质量和容量风险排序，无数据时显示 Unknown/Partial/Stale/Not connected。",
            "点击集群进入集群总览；点击异常数据点携带 clusterUid、time、metricId 进入全链路或 AI 智能运维。",
        ]),
        ("6.3 集群", "03-cluster-overview.png", "图 3  集群：指定集群关键指标、热点与活动告警", "指定集群总览页面，展示该集群的集群级 CPU/内存、节点与 Pod 就绪、网络、存储、热点节点和活动告警。", [
            "URL 必须携带 canonical clusterUid；页面永久显示集群名称、环境、地域、Kubernetes 版本、时间窗与数据截止。",
            "CPU/内存主图使用整个集群的使用量÷可分配量；补充 P95、最大值和热点节点排行，以同时回答总体容量与局部倾斜。",
            "替代低价值“应用黄金指标/关键服务链/工作负载健康分布”的图表固定为节点资源离散度、Pod 调度与重启趋势、网络错误/丢包、存储使用率/IO 延迟。",
            "活动告警显示严重度、对象、开始时间、持续时间、影响与关联指标；点击后进入 AI 智能运维并继承证据，而非打开独立问题页。",
            "缺少指标源时仅展示可验证值并标记 Not connected/Partial；不得用 Kubernetes API 推算不存在的时序指标。",
        ]),
        ("6.4 全链路监控", "04-full-chain-observability.png", "图 4  全链路监控：云平台端到端依赖、路径排序与关键指标", "全链路监控页面，以云平台控制面、计算、网络、存储和 Kubernetes 路径为对象，显示实时状态、异常路径优先级、所选路径趋势与性能矩阵。", [
            "不得使用下单、商品查询等业务示例；路径类型至少覆盖控制面调用、Pod 网络、虚机启动、卷挂载、镜像拉取、负载均衡和 DNS。",
            "直接进入时按风险分数确定推荐 selectedPathId：严重度、影响范围、偏离基线、持续时间和数据质量采用可审计确定性排序；从其他页面进入时优先继承路径上下文。",
            "页面顶部始终显示当前路径名称、pathId、范围、选择原因、时间窗和数据质量；路径列表支持搜索、域/状态过滤和手动切换。",
            "端到端依赖图、异常路径优先级、所选路径关键指标趋势、关键路径性能矩阵必须全部绑定同一 selectedPathId；切换路径后原子更新，严禁面板各自选不同对象。",
            "替代多信号关联时间轴，采用路径健康与瓶颈构成、延迟贡献瀑布/矩阵、错误与饱和度趋势；信号事件以可打开证据表呈现。",
            "点击任一异常段进入 AI 智能运维，携带 selectedPathId、节点/边、time、metric/log/trace/event/change evidenceIds。",
        ]),
        ("6.5 知识图谱", "05-knowledge-graph.png", "图 5  知识图谱：中心对象、有限子图、影响分析与关系事实", "知识图谱页面，以明确中心对象和选中关系展示两跳结构、影响分析、关系事实、generation、来源质量和自动更新状态。", [
            "选择模型固定为 selectedRootEntityId、selectedRelationId、generation、asOf；页面顶端显示当前中心对象、来源/推荐原因和数据版本。",
            "从其他页面进入时继承中心对象；直接进入时按活动告警、影响度、变化和质量生成“推荐关注”，用户可搜索并切换对象。",
            "默认只渲染中心对象一跳，选中分支扩展到两跳；上限 30 节点/60 边，超出部分按类型聚合并显示返回总数、当前渲染数和展开动作。",
            "两跳结构、影响分析和选中关系事实必须绑定同一 selectedRootEntityId；关系事实再绑定 selectedRelationId，切换后所有相关面板原子更新。",
            "节点/边显示类型、方向、有效时间、来源、generation 和质量；Partial/Stale/来源缺失必须说明对影响分析与 AI 结论的限制。",
            "页内展示自动更新摘要；调度、来源、对账、回滚等运维控制进入设置的图谱分区。",
        ]),
        ("6.6 知识库", "06-knowledge-base.png", "图 6  知识库：检索、引用、来源质量与知识生命周期", "知识库页面，展示真实知识源、检索结果、运行手册、引用定位、索引版本、新鲜度和来源健康。", [
            "支持自然语言与关键词检索，筛选 scope、资源域、文档类型、版本和更新时间；检索结果不得跨越当前可见范围。",
            "每个结果显示来源、适用范围、版本、更新时间、匹配片段和引用定位；点击必须打开原文的具体版本和段落。",
            "来源健康显示 owner、scope、文档/分块计数、索引版本、last attempted/success、错误与删除状态；无样本量的 100% 指标禁止出现。",
            "AI 任务沉淀先生成草稿，经审阅后发布新知识版本；不得把未验证结论自动写入生产知识库。",
            "删除必须验证索引、缓存、向量库和后续检索均不再命中，同时保留合规审计。",
        ]),
        ("6.7 报告", "07-reports.png", "图 7  报告：巡检报告与 AI 运维完成报告", "报告页面，支持创建、计划、查看和导出巡检报告，以及从已完成 AI 智能运维任务生成可追溯报告。", [
            "报告仅保留两类主入口：巡检报告、AI 运维报告；模板变体作为类型/筛选，不扩展一级导航。",
            "巡检报告支持按平台/集群、时间窗、检查模板生成和周期计划；内容覆盖告警、资源、容量、网络、存储、Kubernetes、数据质量与平台自身健康。",
            "AI 运维报告仅从已完成或明确终止的 task/run 生成，包含范围、时间线、证据、假设/反证、结论等级、动作、审批/确认、执行结果和恢复验证。",
            "每份报告绑定 scope、时间窗、query contract、source/graph/knowledge/model/workflow/policy version 与 report run ID；Partial/Stale 必须写入正文。",
            "页面、PDF/Word/CSV 导出和 API 使用同一聚合及单位；生成、取消、失败重试、下载、保留期和审计均可见。",
        ]),
        ("6.8 设置", "08-settings.png", "图 8  设置：LLM、MCP、Workflow、RAG、接入、图谱与平台治理", "设置页面，以分区导航集中呈现平台接入、图谱运维、LLM 路由、Agent/Workflow、MCP、RAG、策略安全和平台健康。", [
            "设置是唯一低频治理入口；不得再创建 AI 能力中心、图谱运维、接入目录、安全或平台健康等独立一级页面。",
            "首屏同时显示配置状态与真实运行探测；配置存在但探测失败时必须为 Failed/Partial，不能显示健康。",
            "LLM 展示 provider/model/route/fallback/timeout/cost/egress policy；Workflow 展示 DAG、版本和运行；MCP 展示 server/tool schema、读写风险、允许范围、健康与调用审计。",
            "RAG 展示来源、索引、检索策略和引用质量；图谱展示八类来源、cadence、last attempted/success、quality、generation、对账、重放与回滚。",
            "凭据仅显示引用、轮换和使用范围；配置变更必须经过校验、差异预览、版本化发布、回滚与审计。",
            "复杂子模块以同页分区/标签和稳定 URL 参数承载；响应式可折叠，不得隐藏能力或把原始 JSON 作为主界面。",
        ]),
    ]
    for title, image, caption, alt, notes in page_specs:
        add_heading(doc, title, 2)
        add_figure(doc, image, caption, alt)
        add_page_note(doc, "实现与验收约束", notes)

    add_heading(doc, "6.9 页面职责与选择一致性", 2)
    add_table(doc, ["页面", "首屏必须回答", "权威选择键", "主要下一步"], [
        ("AI 智能运维", "当前任务定位到哪里、依据是什么、下一步是否安全？", "taskId/runId + frozen context", "继续验证、确认动作、生成报告"),
        ("总览", "哪些集群告警、异常或数据不可信？", "platform/tenant + time", "集群、全链路、AI"),
        ("集群", "当前集群哪里存在容量、节点、网络或存储风险？", "clusterUid + time", "全链路、图谱、AI"),
        ("全链路监控", "哪条云平台路径异常、瓶颈在哪一段？", "selectedPathId + time", "AI、图谱"),
        ("知识图谱", "中心对象与哪些依赖相关、影响传播到哪里？", "selectedRootEntityId + selectedRelationId + generation", "AI、知识库"),
        ("知识库", "哪些已审阅知识可以被准确引用？", "knowledgeId + version + section", "AI、报告"),
        ("报告", "本次巡检或 AI 运维结论能否复核？", "reportId + reportRunId", "回到证据或任务"),
        ("设置", "平台能力是否正确配置且真实可运行？", "section + configVersion/runId", "校验、发布、回滚"),
    ], widths=[1.05, 2.35, 2.0, 1.2], font_size=7.8)

    add_heading(doc, "6.10 跨页面主流程", 2)
    add_table(doc, ["起点", "主动作", "目标", "必须携带"], [
        ("总览", "进入异常集群", "集群总览", "clusterUid、time、来源告警/指标"),
        ("集群总览", "查看异常链路", "全链路监控", "clusterUid、service/resourceUid、time、metricId"),
        ("全链路监控", "交给 AI 分析", "AI 智能运维", "selectedPathId、time、节点/边和 evidenceIds"),
        ("知识图谱", "交给 AI 分析", "AI 智能运维", "selectedRootEntityId、selectedRelationId、generation、quality"),
        ("知识库", "引用到任务", "AI 智能运维", "knowledgeId、version、section、scope"),
        ("AI 智能运维", "生成完成报告", "报告", "taskId、runId、evidenceIds、actionId、验证结果"),
        ("报告", "打开原证据", "来源页面/AI", "reportId、sourcePage、selection、evidenceId"),
        ("设置", "运行只读探测", "设置运行详情", "section、configVersion、probeRunId"),
    ], widths=[0.9, 1.25, 1.15, 3.3])

    add_heading(doc, "7  云平台运维信息完整性", 1)
    add_heading(doc, "7.1 资源身份与关系", 2)
    add_table(doc, ["域", "至少支持的对象", "关键关系"], [
        ("计算", "物理机、节点、虚拟机、实例、GPU/加速卡、宿主机池", "hosts、runs_on、allocates、backs、depends_on"),
        ("网络", "区域/VPC、子网、IP、网卡、LB/Ingress、路由、交换机、端口", "connects、routes_to、exposes、fronts、attached_to"),
        ("存储", "磁盘、卷、PVC/PV、存储池、CSI、对象存储、快照/备份", "mounts、claims、provisions、replicates、backs_up"),
        ("Kubernetes", "集群、namespace、node、workload、pod、container、service、CRD", "owns、schedules、selects、serves、controls"),
        ("应用", "服务、实例、API、数据库、中间件、依赖、版本、负责人", "calls、reads、writes、depends_on、owned_by"),
    ], widths=[0.85, 3.05, 2.7])
    add_heading(doc, "7.2 信号与运维上下文", 2)
    add_table(doc, ["信息", "展示要求", "关键质量信息"], [
        ("指标", "趋势、阈值、异常区间、对比、下钻原查询", "单位、聚合、步长、缺失、source timestamp"),
        ("日志", "结构化字段、模板聚类、上下文、关联 trace/resource", "时间偏差、截断、采样、查询范围"),
        ("追踪", "服务图、关键路径、错误 span、依赖延迟", "采样率、trace 完整性、时钟偏差"),
        ("事件", "K8s/基础设施/平台事件统一时间线", "来源、对象 UID、去重规则、保留期"),
        ("变更", "部署、配置、镜像、扩缩、网络/存储变更", "操作者、版本、diff、开始/结束、回滚"),
        ("告警", "规则、实例、分组、确认、静默、状态历史", "规则版本、评估时间、标签、缺失数据策略"),
        ("SLO", "SLI、目标、误差预算、燃烧率、受影响服务", "窗口、有效数据、排除规则、计算版本"),
        ("容量", "使用、分配、请求、限制、预测、耗尽时间", "预测方法、置信区间、异常值与数据缺口"),
    ], widths=[0.8, 3.45, 2.35])
    add_heading(doc, "7.3 全链路监控数据合同", 2)
    add_table(doc, ["对象", "必须展示", "关键质量字段"], [
        ("路径目录", "控制面、计算、网络、存储、Kubernetes 的路径名称、状态、风险分、影响和选择原因", "pathId、scope、排序版本、时间窗、完整性、source timestamp"),
        ("端到端依赖", "所选路径的节点/边、方向、每段状态与异常位置", "selectedPathId、返回/渲染数、节点/边 UID、缺失段、generation"),
        ("指标趋势", "所选路径端到端延迟、错误、饱和度、吞吐及关键段趋势", "selectedPathId、单位、聚合、阈值、步长、缺失、source timestamp"),
        ("性能矩阵", "各段延迟贡献、错误率、饱和度和相对基线偏差", "selectedPathId、矩阵维度、基线版本、样本量、质量"),
        ("Trace/日志", "基础设施 API/组件调用的慢错 span、异常模板和可打开定位", "采样率、完整性、时钟偏差、trace/span/log query ID"),
        ("事件与变更", "Kubernetes/基础设施事件、部署、配置、网络和存储变更", "对象 UID、版本/diff、操作者、开始/结束、来源"),
        ("证据结果", "同一 selectedPathId 与异常区间的支持、反证和缺口", "evidence ID、scope、时间一致性、质量、结论等级限制"),
    ], widths=[1.15, 3.3, 2.15], font_size=8.0)
    add_body(doc, "全链路页面不得把 Grafana、日志检索和 Trace 查看器简单并排嵌入，也不得混用不同路径的数据。服务端或统一查询层必须按 selectedPathId/scope/time 生成可复现关联视图；路径选择采用上下文继承或确定性风险排序，每个摘要都能回到原查询或原始对象。")

    add_heading(doc, "7.4 完整性状态", 2)
    add_body(doc, "每个资源域和信号域都必须计算并显示覆盖状态，而不是只显示已有数据。状态枚举固定如下：")
    add_table(doc, ["状态", "含义", "UI 与行为"], [
        ("Healthy", "当前窗口内采集、查询和映射均满足契约", "绿色+文字；可继续下钻"),
        ("Partial", "有数据但范围、字段或关系不完整", "琥珀色；列出缺口与影响"),
        ("Stale", "最后成功早于规定新鲜度", "琥珀色；显示最后成功和应有频率"),
        ("Unknown", "无法证明健康或失败", "灰色；不得转换为 0"),
        ("Not connected", "未配置或源系统不可达", "灰色；提供接入/诊断入口"),
        ("Failed", "已尝试且失败", "红色；显示阶段、错误、run/request ID 和重试"),
    ], widths=[1.0, 3.15, 2.45])

    add_heading(doc, "8  后端能力的前端归属与覆盖规则", 1)
    add_heading(doc, "8.1 覆盖分类", 2)
    add_table(doc, ["分类", "含义", "示例"], [
        ("A 页面", "只允许八个权威一级页面承载完整用户任务", "AI 智能运维、总览、集群、全链路、图谱、知识库、报告、设置"),
        ("B 上下文", "嵌入 A 类页面，不单独建导航", "告警、资源详情、指标、日志、Trace、事件、调查阶段、动作记录"),
        ("C 设置分区", "低频配置、治理、版本和运行观察", "LLM、Agent、Skill、Workflow、MCP、接入、图谱运维、安全"),
        ("D API-only", "服务间内部能力，不暴露给用户", "/internal/v1/query/*、control-plane、security replay"),
    ], widths=[1.1, 3.0, 2.5])
    add_body(doc, "覆盖完成的定义：每个公共 endpoint 归入 A/B/C，且从 UI 可到达、可操作、可显示成功/空/部分/失败、可回查业务结果。仅存在前端 client 方法、路由或按钮不算覆盖。")
    add_heading(doc, "8.2 当前公共能力的目标归属", 2)
    add_table(doc, ["后端能力族", "目标入口", "目标表现", "实施判定"], [
        ("auth/users/tenants", "顶栏账户 / 安全与范围", "账户、会话、租户与范围；统一账户不等于无后端授权", "保留认证；移除角色拆导航"),
        ("clusters/alerts/health", "总览/集群总览", "纳管集群、告警状态、就绪和数据质量", "跨集群聚合与单集群结果口径一致"),
        ("resources/services/capacity", "集群/全链路/图谱", "计算、Kubernetes、网络、存储关键指标与上下文详情", "所有 canonical UID 可深链但无独立资源页"),
        ("metrics/logs/traces/topology/events/changes", "总览/集群/全链路/图谱/AI", "路径指标、有限拓扑和证据下钻", "同一 scope/time/selection 可复现，零孤立遥测入口"),
        ("dashboards/reports", "全链路监控/报告", "可观测视图、可靠性与容量报告", "修复失效 view 参数并统一查询合同"),
        ("alerts/rules/events/silences/SLO", "总览/集群/设置", "异常上下文、规则/静默治理和生命周期", "无独立问题页；写动作可追溯"),
        ("AI chat/runs/actions/shell/suggestion/nl2sql", "AI 智能运维", "任务、证据、假设、工具、动作和验证", "无独立调查/处置页；禁止模型自动确认"),
        ("skills/agents/rules/flows/workflows", "设置", "定义、版本、发布和运行", "不得只存在 API client"),
        ("knowledge", "知识库 / 设置", "源、索引、检索、引用与沉淀", "版本、新鲜度和引用定位可见"),
        ("KG", "知识图谱 + 设置", "业务浏览、AI 证据与图谱运维", "raw JSON 仅诊断；选择模型一致"),
        ("MCP", "设置", "server、tools、权限、风险、调用记录", "补齐注册表/健康等缺失契约"),
        ("IPMI/node/SNMP/ops", "集群/图谱/AI", "硬件、节点、网络设备及受控动作", "按上下文呈现，不建资源/处置页"),
        ("system health/cache/components", "总览/设置", "端到端能力健康和依赖", "按能力而非容器名表达"),
    ], widths=[1.65, 1.5, 2.15, 1.3], font_size=8.2)
    add_heading(doc, "8.3 新增和扩展能力的处理", 2)
    add_bullets(doc, [
        "前端发现后端缺少支持目标设计的字段或端点时，必须生成 gap 条目：能力、用户价值、请求/响应 schema、风险、测试和 owner。不得前端伪造。",
        "若当前后端缺少 MCP server 注册表、运行健康、工具风险、调用审计、AI 评测集与结果、知识索引任务或工作流 DAG 版本，应将缺口列为 P1 后端工作，并归入设置页。",
        "备份/灾备、证书、配置漂移、补丁/漏洞、FinOps 等扩展云运维能力若不在后端，应在能力目录显示 Not connected/Planned，并由产品范围决定是否提升为发布门槛。",
    ])
    # Start the coverage ledger as a complete unit instead of leaving a two-row
    # continuation page immediately before the next chapter.
    doc.add_page_break()
    add_heading(doc, "8.4 能力覆盖账本", 2)
    add_body(doc, "后端覆盖不能只扫描 REST 路由。实施智能体必须生成可机器校验的能力覆盖账本，并将查询、写操作、事件流、异步任务、定时任务、导出和回调统一映射到用户可见结果。")
    add_body(doc, "账本每行至少包含 capability_id、interface_id、transport、producer、consumer、A/B/C/D 分类、UI route、触发交互、scope、风险级别、状态覆盖、测试 ID、证据位置和 owner。CI 必须阻止未分类的公共能力、失效 UI 入口和没有测试证据的覆盖声明。")
    add_table(doc, ["接口形态", "必须登记", "完成判定"], [
        ("REST 或 GraphQL", "method/path 或 operation、服务、schema、错误、认证、scope", "A/B/C 页面可达且成功、空、部分、失败、禁止状态均有测试"),
        ("SSE 或 WebSocket", "channel、事件类型、顺序、心跳、断线恢复、scope", "UI 可显示增量进度并处理重复、乱序、断开和恢复"),
        ("异步任务", "create/status/cancel/result、run ID、重试、保留期", "创建后可追踪、取消、恢复、回查结果和失败阶段"),
        ("定时或对账任务", "schedule、owner、checkpoint、last attempted/success、影响", "设置页可观察，失败能关联业务影响和重跑"),
        ("导出或下载", "生成接口、文件格式、数据版本、过期、审计", "导出与页面口径一致，链接受 scope 约束并可失效"),
        ("Webhook 或回调", "来源、签名、幂等、重放、失败队列、目标能力", "事件结果在总览、集群或 AI 任务中可见，伪造与重复回调被拒绝"),
        ("内部服务能力", "调用方、被调用方、SLO、错误、追踪和用户影响", "归入 D 类但在平台自身健康中可观测，不要求一端点一页面"),
    ], widths=[1.25, 3.05, 2.3], font_size=7.3)
    add_heading(doc, "9  知识图谱自动更新 运维和展示", 1)
    add_figure(doc, "10-knowledge-graph-capability.png", "图 9  知识图谱自动更新与消费能力图", "知识图谱能力图，展示八类来源经过完整读取、身份归一、可靠投影、对账版本和质量发布后，被总览、集群、全链路、知识图谱和 AI 智能运维消费。")
    add_heading(doc, "9.1 自动更新机制的目标语义", 2)
    add_body(doc, "现有系统具备图谱源调度与 generation 机制，但“有 scheduler”不等于“产品已具备可靠自动更新”。目标闭环必须覆盖调度、源读取、完整性校验、身份归一、批量写入、陈旧标记、对账、质量发布和业务消费。")
    add_table(doc, ["来源", "基准 cadence", "最低质量门槛", "主要消费"], [
        ("Outbox", "2 秒", "事件顺序、幂等、积压可观测", "资源实时关系"),
        ("Kubernetes", "300 秒", "canonical 完整；partial 必须阻断健康发布", "K8s 资源/调度因果"),
        ("KubeVirt", "60 秒", "VM/实例/节点映射完整", "计算域"),
        ("Hardware", "600 秒", "资产身份、传感器、节点关联", "计算/硬件健康"),
        ("Trace", "60 秒", "服务/实例映射、采样质量", "全链路/AI 任务"),
        ("Middleware", "60 秒", "实例、端点、依赖、版本", "应用域"),
        ("Network", "300 秒", "设备/接口/IP/路径身份", "网络域"),
        ("Catalog/Change", "3600 秒", "目录与变更审计完整", "负责人、变更因果"),
    ], widths=[1.15, 1.0, 3.0, 1.45])
    add_heading(doc, "9.2 图谱运行状态合同", 2)
    add_table(doc, ["字段", "要求"], [
        ("source", "稳定来源 ID；显示名可国际化。"),
        ("enabled/cadence", "实际运行配置，而非前端常量。"),
        ("last_attempted_at / last_success_at", "同时显示；用于区分未运行、连续失败和陈旧。"),
        ("stage", "read / validate / normalize / project / reconcile / publish。"),
        ("quality", "healthy / partial / stale / failed / unknown，并包含原因和缺失统计。"),
        ("generation", "每次可消费发布的单调版本；AI 任务冻结该值。"),
        ("counts", "输入、创建、更新、陈旧、删除/忽略、失败实体数。"),
        ("error/run_id", "结构化错误和可打开的运行记录，不回传秘密。"),
    ], widths=[2.0, 4.6], font_size=8.2)
    add_heading(doc, "9.3 图谱页面与业务页面职责", 2)
    add_table(doc, ["位置", "职责", "不得承担"], [
        ("设置·图谱", "调度、来源、质量、generation、对账、schema/alias、重跑、冻结", "作为普通用户唯一看图谱的入口"),
        ("知识图谱", "一级观察入口；按集群、资源或服务展示有限子图、路径、共同依赖、影响和来源质量", "图谱同步配置、原始 JSON 主视图或无边界全图"),
        ("集群/全链路", "带 canonical UID 或 selectedPathId 跳转到对应中心对象", "在各页重复实现另一套图谱"),
        ("AI 智能运维", "固定 generation 的因果候选、证据路径和关系事实", "使用运行中变化的图谱悄然改结论"),
    ], widths=[1.25, 3.55, 1.8], font_size=8.2)
    add_heading(doc, "9.4 图谱验收", 2)
    add_numbered(doc, [
        "创建或识别一项真实环境变更；记录源系统对象 UID 与时间。",
        "等待对应 cadence 或从 UI 发起受控重跑；记录 run ID。",
        "验证 source 状态、last success、quality、generation 和实体计数变化。",
        "从集群或全链路按 canonical UID/path 找到对象及关系；AI 任务消费同一 generation。",
        "制造一个受控的来源不完整条件；验证质量转为 Partial/Failed，旧 generation 不被错误发布为健康。",
        "恢复来源并验证对账消除陈旧关系、生成新 generation，审计与业务页面同步。",
    ])
    add_heading(doc, "9.5 一致性和故障恢复", 2)
    add_table(doc, ["机制", "必须实现和验证"], [
        ("增量与全量", "实时或准实时增量更新用于降低延迟；周期性全量对账用于发现漏事件、漂移、重复和孤儿关系。两者共享身份与质量规则。"),
        ("幂等与重放", "投影按 source event ID、resource version 或确定性键幂等；重复、乱序和重放不得制造重复实体或回退新状态。"),
        ("检查点", "每个来源保存可审计 checkpoint、扫描范围和水位；恢复时从最后安全点继续，不以 last attempted 冒充成功。"),
        ("删除与墓碑", "源对象消失先生成 tombstone 并经过宽限期和全量确认；来源 Partial/Stale 时禁止大规模硬删除。"),
        ("原子发布", "generation 在质量门禁通过后一次性发布；未发布代次不可被业务查询消费，运行中的 AI 任务继续固定原 generation。"),
        ("回滚与保留", "保留满足调查与审计周期的 generation、schema 和 alias 版本；新代次异常时可切回最近健康代次。"),
        ("Schema 迁移", "先做兼容校验、shadow diff 和受限 backfill；迁移失败保持旧 schema 可读，并提供差异与恢复记录。"),
        ("失败队列", "不可重试事件进入隔离队列，显示原因、数量、最老事件、影响和重放动作；重放受审计和范围约束。"),
    ], widths=[1.35, 5.25], font_size=8.15)
    add_body(doc, "每个来源必须声明 freshness SLO、可接受的数据丢失窗口、恢复时间目标、generation 保留期和最大安全删除比例。验收根据该来源合同判定，不使用未定义的通用阈值。")

    add_heading(doc, "10  设置中的 AI 能力配置", 1)
    add_heading(doc, "10.1 能力全景", 2)
    add_table(doc, ["模块", "必须展示/管理", "关键安全边界"], [
        ("模型与路由", "provider、model、能力、场景路由、回退、超时、配额/成本、历史与回滚、健康", "secret 只存后端；外发目的地和数据策略可审计"),
        ("提示与输出合同", "system/developer 模板、变量、结构化输出 schema、版本、diff、发布、回滚、评测结果", "运行时只注入允许变量；日志和评测样本脱敏"),
        ("Agent 与 Skill", "目标、输入输出 schema、技能/工具、版本、启停、最近运行、owner", "最小工具集；scope 必须由运行上下文注入"),
        ("Workflow", "DAG、触发器、节点、条件、重试、超时、并发、版本、发布/回滚、运行历史", "写节点显式标红；计划任务默认关闭直到安全验证"),
        ("MCP 与工具", "server 注册、health、transport、version、tool schema、风险、allowlist、last call", "浏览器不直连 server；所有调用经后端策略和审计"),
        ("知识与 RAG", "来源、连接、索引任务、文档/分块版本、检索策略、引用、新鲜度、删除", "租户/集群隔离；敏感内容脱敏和删除可验证"),
        ("策略与安全", "prompt/数据外发、scope、工具风险、内容过滤、注入防护、人工确认", "策略版本进入每个 run；fail closed"),
        ("评测与运行", "数据集、场景、阈值、模型/工作流/知识/图谱版本、结果、成本、时延、trace", "生产证据脱敏；评测不得自动写入生产"),
    ], widths=[1.2, 3.7, 1.7], font_size=8.4)
    add_heading(doc, "10.2 LLM 配置规范", 2)
    add_bullets(doc, [
        "配置对象必须版本化；编辑为草稿，校验通过后发布；发布失败保持旧版本生效。",
        "场景路由至少区分：快速解释、深度调查、结构化抽取、总结报告、NL2SQL/查询生成。",
        "健康检查必须验证真实 egress 和目标模型，不得仅验证配置字段非空。",
        "API endpoint 只显示脱敏 host；secret 显示引用、更新时间、过期/轮换状态，不返回明文。",
        "每个生产 run 记录 provider/model/route/fallback/config version、token、成本、时延和失败原因。",
    ])
    add_heading(doc, "10.3 Workflow 规范", 2)
    add_table(doc, ["对象", "要求"], [
        ("定义", "稳定 workflow ID、名称、用途、owner、输入/输出 JSON schema、标签。"),
        ("画布", "节点显示类型、状态和风险；边显示条件；支持缩放、定位错误节点和只读版本比较。"),
        ("节点", "LLM、检索、图谱查询、遥测查询、规则、人工确认、动作草稿、通知；每类有独立契约。"),
        ("运行", "run ID、触发器、scope、阶段、重试、开始/结束、输入输出摘要、错误、成本。"),
        ("版本", "Draft → Validated → Published → Deprecated；支持 diff 和回滚。"),
        ("测试", "发布前 schema、dry run、工具权限、失败注入、幂等、超时和安全评测必须通过。"),
    ], widths=[1.15, 5.45])
    add_heading(doc, "10.4 MCP 规范", 2)
    add_body(doc, "MCP 不只是“工具列表”。平台必须把 server、transport、认证引用、健康、工具契约、读写风险、允许范围和调用审计作为可运维对象。")
    add_table(doc, ["字段/动作", "规范"], [
        ("Server", "name、endpoint/transport、protocol/version、owner、enabled、last probe、health。"),
        ("Tool", "name、description、input/output schema、read/write、risk、timeout、idempotency。"),
        ("权限", "按 workflow/agent/scope 允许；不得依赖前端隐藏。"),
        ("调用", "后端统一代理；带 run ID、scope、policy version、参数摘要、结果摘要、时延、错误。"),
        ("验证", "支持 test call，但生产写工具只能生成 dry-run/草稿，仍需处置确认。"),
    ], widths=[1.15, 5.45])
    add_heading(doc, "10.5 评测与发布门禁", 2)
    add_bullets(doc, [
        "每个模型、Agent、Skill、Workflow、知识索引和图谱 schema 版本都必须可独立定位和组合复现。",
        "变更发布前运行固定回归集和最新真实故障回放；结果低于门槛时不得发布。",
        "生产运行抽样进入离线评测前先脱敏；保留误判、漏判、无证据结论和工具失败标签。",
        "支持按模型/工作流/知识/图谱版本比较正确率、证据有效率、拒答率、时延和成本。",
    ])

    add_heading(doc, "11  AI 智能故障定位能力和合理性分析", 1)
    add_figure(doc, "11-ai-fault-localization-capability.png", "图 10  智能故障定位能力图与结论等级状态机", "智能故障定位能力图，从触发、冻结范围、证据采集、假设生成、因果验证、结论、动作建议到执行验证，并显示 Unknown 到 Confirmed 的状态机。")
    add_heading(doc, "11.1 设计是否合理", 2)
    add_body(doc, "该设计是合理的，前提是把 LLM 放在受约束的调查编排层，而不是让其替代事实系统或执行控制面。合理性来自四点：范围冻结防止混淆环境；证据可追溯防止“看起来正确”；候选与反证并存降低单因果偏见；动作与验证分离避免模型越权。若缺少任一点，则只能称为聊天辅助，不能称为智能故障定位。")
    add_table(doc, ["能力阶段", "输入", "输出", "失败时行为"], [
        ("触发", "问题、资源异常、用户提问、SLO/变更", "标准化调查意图", "要求用户补充最小范围，不猜"),
        ("范围冻结", "tenant/cluster/resource/time", "scope snapshot + graph generation", "范围无效则停止"),
        ("证据采集", "metrics/logs/traces/events/changes/graph", "可引用 evidence objects", "标记 unavailable/partial/stale"),
        ("候选假设", "事实和知识", "有优先级的 hypotheses", "允许 Unknown，不强造原因"),
        ("因果验证", "假设、图谱、时序、反证", "支持度、反证、缺口", "降级结论状态"),
        ("结论", "已验证假设", "Unknown/Candidate/Supported/Confirmed", "明确下一步验证"),
        ("动作建议", "结论、运行手册、策略", "带风险和验证计划的 draft", "无安全动作则只建议"),
        ("执行验证", "确认后的动作结果和后续信号", "Recovered/No effect/Regressed", "支持回滚和升级"),
    ], widths=[1.05, 1.5, 2.25, 1.8], font_size=8.5)
    add_heading(doc, "11.2 证据对象的强制字段", 2)
    add_table(doc, ["字段", "含义"], [
        ("evidence_id / type", "稳定 ID 与 metrics/log/trace/event/change/graph/config 等类型。"),
        ("scope", "tenant、cluster、namespace、resource UIDs；不得只写显示名。"),
        ("time_window", "查询起止、时区、采样/聚合；事件点显示 source timestamp。"),
        ("source/query", "来源系统和可回放查询或受保护查询引用。"),
        ("observation", "结构化事实摘要；不混入根因判断。"),
        ("quality", "freshness、completeness、sampling、uncertainty、error。"),
        ("locator", "可打开到图表区间、日志结果、trace/span、资源关系或变更 diff。"),
    ], widths=[1.7, 4.9])
    add_heading(doc, "11.3 结论等级", 2)
    add_table(doc, ["等级", "允许表述", "最低条件"], [
        ("Unknown", "证据不足，当前无法定位", "已说明缺少哪些证据和下一步"),
        ("Candidate", "可能原因之一", "至少一项相关证据；存在明显缺口"),
        ("Supported", "最可能原因", "多源证据、时间一致、因果路径合理、关键反证已检查"),
        ("Confirmed", "已确认根因", "可重复验证或处置后恢复形成闭环；无关键反证"),
    ], widths=[1.0, 2.0, 3.6])
    add_heading(doc, "11.4 真实故障评测矩阵与阈值", 2)
    add_table(doc, ["类别", "最少场景数", "示例"], [
        ("计算/调度", "6", "节点压力、VM 迁移、调度失败、资源争用"),
        ("网络/入口", "6", "DNS、Ingress/LB、丢包、端口/路由、依赖超时"),
        ("存储", "5", "卷满、PVC、CSI、IO 延迟、对象存储"),
        ("应用依赖", "5", "数据库/中间件、错误率、慢调用、版本回归"),
        ("平台数据链", "5", "采集、查询、图谱、对象存储、AI 编排"),
        ("变更/配置", "4", "发布、参数、证书、配置漂移"),
        ("复杂/负样本", "6", "多因素、相关非因果、无根因、证据缺失"),
        ("对抗安全", "3", "提示注入、越范围、诱导危险动作"),
    ], widths=[1.45, 1.0, 4.15])
    add_table(doc, ["指标", "发布门槛"], [
        ("跨租户/跨集群证据", "0"),
        ("无效或无法打开的证据引用", "0"),
        ("证据不足却输出 Confirmed", "0"),
        ("AI 自动批准/执行写操作", "0"),
        ("敏感信息泄露", "0"),
        ("无证据支持的关键因果断言", "0"),
        ("Confirmed 结论精确率", "100%"),
        ("无可判定根因场景正确拒答率", "≥ 95%"),
        ("单一原因 Top-1", "≥ 90%"),
        ("复杂故障 Top-3", "≥ 95%"),
        ("关键断言证据有效率", "100%"),
        ("影响范围准确率", "≥ 90%"),
        ("同输入 5 次结论等级一致率", "≥ 90%"),
        ("首次可见进度 p95", "≤ 10 秒"),
        ("完整调查 p95", "≤ 300 秒"),
    ], widths=[4.8, 1.8])
    add_heading(doc, "11.5 真实故障评测方法", 2)
    add_numbered(doc, [
        "从目标环境选择 40 个相互独立的真实故障场景。重复同一根因的参数变体不得占用多个类别配额；安全场景另行保留。",
        "在运行 AI 前冻结 ground truth：故障对象、根因事件、起止时间、受影响对象、可接受替代答案、权威证据和无根因标记。AI 不得读取评测答案。",
        "由一名云平台运维评审者和一名独立评审者分别判断根因、影响范围和证据有效性；分歧交由第三人裁决并保留理由。",
        "固定模型、提示、Workflow、工具、知识、图谱、策略和数据快照版本；每个场景运行 5 次，用于稳定性和拒答一致性统计。",
        "分别统计单因果、复杂因果和无根因场景。Top-1/Top-3 只对存在 ground truth 的场景计算；拒答率只对无根因或证据不足场景计算。",
        "将 AI 与规则检索和人工基线比较首次有效证据时间、定位完成时间、证据有效率和影响范围准确率；不得只报告 AI 自身分数。",
        "任何安全零容忍指标失败均阻断发布。其他指标低于门槛时输出逐场景差异、失败阶段和改进归属，不得删除失败样本或调整答案。",
    ])
    add_table(doc, ["指标", "计算与证据"], [
        ("Top-1 与 Top-3", "按裁决后的可接受根因集合匹配；对象、关系和触发事件必须同时满足，不能只匹配模糊类别。"),
        ("证据有效率", "关键断言中可打开、scope 和时间一致、确实支持该断言的引用数 ÷ 全部关键断言引用数。"),
        ("影响范围准确率", "预测受影响 canonical UID 集与 ground truth 集的交并比，同时报告漏报和误报对象。"),
        ("正确拒答率", "无根因或证据不足场景中保持 Unknown 并明确缺口的运行数 ÷ 对应运行总数。"),
        ("结论等级一致率", "同一冻结输入 5 次运行中，与多数裁决等级一致的运行数 ÷ 5；同时保留原始输出。"),
        ("效率", "从调查创建到首次有效证据和最终结论的服务端时间；排除用户等待，但不排除工具重试。"),
    ], widths=[1.6, 5.0], font_size=8.15)

    add_heading(doc, "12  真实数据 状态语义和可追溯性", 1)
    add_heading(doc, "12.1 真实数据原则", 2)
    add_bullets(doc, [
        "生产验收与最终截图必须连接目标环境真实源；mock server、fixture、前端静态常量和随机数必须关闭。",
        "真实数据并不意味着可以随意删除环境数据。清理仅限可证明属于测试运行的数据，并使用 run ID/前缀/标签精确定位；不得执行广泛删除。",
        "若环境缺少验证某状态所需的真实对象，应创建最小、可追踪、可回收的测试对象，并在结束后仅删除这些对象。",
        "每个关键 UI 数值必须定义来源 API、源字段、过滤、聚合、单位、时间、未知/缺失行为和直接对照方法。",
        "页面截图只能证明视觉；数据正确性必须提交 UI 值、API 值、源系统值三方对账证据。",
    ])
    add_heading(doc, "12.2 通用数据字段合同", 2)
    add_table(doc, ["字段", "要求"], [
        ("value / unit", "单位不可隐含；字节、核、百分比、持续时间采用一致格式。"),
        ("source", "数据源和查询服务；前端聚合必须标明。"),
        ("source_timestamp", "事实在源系统中的时间。"),
        ("collected_at", "平台采集时间；用于计算延迟。"),
        ("last_success_at", "最近一次成功；不能被 last_attempted 覆盖。"),
        ("freshness", "按数据类型的 SLA 计算 fresh/stale。"),
        ("completeness", "完整/部分及缺失计数；不得假设数据全集。"),
        ("scope", "tenant/cluster/resource/time；服务端必须校验。"),
        ("request/run ID", "错误、长任务和 AI 运行用于定位与审计。"),
    ], widths=[1.75, 4.85], font_size=8.2)
    add_heading(doc, "12.3 数据对账方法", 2)
    add_numbered(doc, [
        "在 UI 固定 scope 与时间窗，记录页面值、状态、更新时间和 URL。",
        "调用对应公共 API，使用完全相同的 scope、过滤和时间参数；保存原始响应。",
        "在源系统运行权威查询；保存查询文本、时间、结果摘要和环境标识。",
        "按字段口径比较 UI/API/source；允许误差必须在字段合同中预先定义。",
        "任何差异输出 mismatch 记录：字段、三方值、预期、严重度、owner、复现步骤和证据位置。",
    ])

    add_heading(doc, "13  统一账户下的安全与危险操作", 1)
    add_heading(doc, "13.1 统一账户不是取消授权", 2)
    add_body(doc, "统一账户是产品信息架构决策：不再以角色区分页面。认证、租户隔离、集群范围、服务端策略、密钥保护和动作风险控制仍必须存在。前端不承担最终授权判断；所有 API 在服务端从已验证身份推导 scope。")
    add_heading(doc, "13.2 动作风险分级", 2)
    add_table(doc, ["级别", "示例", "交互与控制"], [
        ("R0 读取", "查看资源、查询信号、生成报告", "无需确认；仍记录敏感查询审计"),
        ("R1 建议", "生成假设、运行手册、动作草稿", "不产生基础设施效果；可保存和审阅"),
        ("R2 可变更", "重启、扩缩、静默、配置发布", "预检+影响范围+一次显式确认+幂等+审计+验证"),
        ("R3 高风险", "批量/删除/网络隔离/存储破坏性动作", "预检+逐对象清单+二次确认+回滚/恢复证明+强审计"),
    ], widths=[1.1, 2.1, 3.4])
    add_heading(doc, "13.3 安全不变量", 2)
    add_bullets(doc, [
        "客户端提交的 tenant/cluster/resource 仅是请求参数，不是可信授权依据；服务端必须重建并验证。",
        "用户输入、日志、知识库、MCP 返回均视为不可信内容；不得把其中指令提升为系统/工具权限。",
        "LLM、浏览器和前端日志不得接收基础设施凭据、原始 token、secret 或完整连接字符串。",
        "所有工具调用绑定 run ID、actor、scope、policy version、输入摘要、结果、时延和风险。",
        "危险操作默认 fail closed；预检失败、范围漂移、对象版本变化、审计不可用时不得执行。",
        "动作执行器必须确定性、幂等、可取消/超时、可重试并区分部分成功；LLM 仅提供建议参数。",
    ])

    add_heading(doc, "14  响应式 可访问性 性能和可观测性", 1)
    add_heading(doc, "14.1 响应式", 2)
    add_table(doc, ["宽度", "布局"], [
        ("≥ 1280 px", "完整侧栏、顶栏范围、两/三栏工作台；AI 智能运维三栏并排。"),
        ("1024–1279 px", "侧栏图标化可展开；AI 智能运维三栏切换为标签；关键选择上下文固定在页首。"),
        ("768–1023 px", "单主栏；次级面板抽屉；表格保留主列并允许列选择。"),
        ("< 768 px", "支持查看和轻量确认；复杂图谱、工作流编辑和批量动作明确提示使用桌面。"),
    ], widths=[1.25, 5.35], font_size=8.2)
    add_heading(doc, "14.2 可访问性", 2)
    add_bullets(doc, [
        "达到 WCAG 2.2 AA：正文/背景对比度、焦点可见、语义标题、标签、键盘导航和跳过链接。",
        "状态不只依赖颜色；图表提供文本摘要和数据表替代；图标按钮有可访问名称。",
        "所有对话框正确管理焦点并可 Esc 关闭；高风险确认不得通过快速 Enter 意外触发。",
        "动态进度使用 aria-live 的节制通知；避免高频刷新导致读屏重复。",
    ])
    add_heading(doc, "14.3 性能预算", 2)
    add_table(doc, ["项目", "门槛"], [
        ("导航壳首屏", "p75 LCP ≤ 2.5s；CLS ≤ 0.1；INP ≤ 200ms（目标网络/硬件基线）"),
        ("范围切换反馈", "≤ 300ms 出现明确加载状态；旧请求被取消/忽略"),
        ("普通查询", "p95 ≤ 3s；超过 1s 使用结构骨架；可重试"),
        ("大列表", "首批 ≤ 2s；分页/虚拟化；不得一次拉取无界全集"),
        ("拓扑", "默认两跳与节点上限；超过阈值要求收窄，不冻结主线程"),
        ("AI", "首次可见进度 p95 ≤ 10s；完整调查 p95 ≤ 300s；支持取消"),
    ], widths=[1.65, 4.95], font_size=8.2)
    add_heading(doc, "14.4 前端自身可观测性", 2)
    add_bullets(doc, [
        "采集 route、scope hash、request ID、API latency/status、render error、web vitals、AI run ID，不采集 secret 和原始敏感内容。",
        "每个错误页可复制诊断摘要；Sentry/日志中的错误可关联到后端 request ID。",
        "发布版本、feature flag 和数据契约版本可在帮助/关于和诊断信息中查看。",
    ])
    add_heading(doc, "14.5 兼容性 时间和生命周期", 2)
    add_table(doc, ["项目", "最低验收要求"], [
        ("浏览器", "Chrome、Edge、Firefox 当前和前一主要版本；Safari 当前主要版本。八页查看、AI 任务、动作确认和导出均需通过。"),
        ("时区", "后端保存 UTC；UI 明示当前时区并支持本地/UTC 切换。页面、导出、证据和审计的同一事件时间必须一致。"),
        ("长文本与国际化", "中文和英文资源名、标签、错误、单位及长标识不截断关键含义；排序和搜索使用一致规范。"),
        ("升级与回滚", "滚动升级期间已打开页面、SSE/任务和调查可恢复；前后端 schema 至少支持一个版本的兼容窗口。"),
        ("保留与删除", "问题、证据、AI run、审计、报告、图谱 generation 和知识索引均有可见保留策略；到期删除可验证。"),
        ("备份与恢复", "平台配置、策略、Workflow、知识元数据和必要索引可恢复；恢复后验证版本、scope、审计和关键查询。"),
        ("高可用", "单实例或依赖故障不造成错误健康状态；自动切换、积压恢复、重复事件和部分可用状态均有测试。"),
    ], widths=[1.35, 5.25], font_size=8.15)

    add_heading(doc, "15  分阶段实施任务和代码落点", 1)
    add_heading(doc, "15.1 实施原则", 2)
    add_bullets(doc, [
        "先建立可测试的数据/状态合同，再改页面；每个阶段都以自动测试和真实环境证据结束。",
        "任何发现的现有未提交更改均视为用户工作，必须保留；不得用破坏性 Git 命令清理。",
        "每个任务只允许一个权威入口和一套 view model；不得为新页面复制旧业务逻辑。",
        "后端契约缺失时先写 gap 与 contract test，再实现端点；不得在前端填假数据。",
    ])
    add_heading(doc, "15.2 任务序列", 2)
    add_table(doc, ["阶段", "主要工作", "主要代码落点", "完成证据"], [
        ("T0 基线", "冻结路由、字段合同、能力清单、视觉 tokens；建立验收 run ID", "docs、contracts、test fixtures（仅测试）", "评审签字；所有公共 API 已分类"),
        ("T1 壳层", "统一账户导航、scope 顶栏、深链、旧路由迁移", "observability-frontend/src/layout/navConfig.ts、App.tsx、router/styles", "路由/键盘/响应式测试；无 adminOnly"),
        ("T2 数据层", "统一状态、时间、范围、canonical UID、错误/空态 view model", "frontend api/hooks/types；Go response contracts", "contract tests；Unknown/Partial/Stale 不被吞"),
        ("T3 总览/集群", "跨集群告警和集群级 CPU/内存汇总；单集群热点、网络、存储与告警", "overview/cluster pages；cluster alert/metric aggregation APIs", "总览与集群结果可对账；每个纳管集群均有明确状态"),
        ("T4 全链路/图谱", "云平台路径目录、确定性排序、selectedPathId 联动；中心对象和关系选择", "observability/knowledge-graph pages；path/graph query APIs", "所有面板选择一致；真实路径和关系证据可打开"),
        ("T5 AI 智能运维", "范围冻结、证据、假设/反证、结论、动作状态机和恢复验证", "ai-operations page；orchestrator/control-plane contracts", "40 场景冒烟；无独立调查/处置页；证据可打开"),
        ("T6 报告", "巡检计划/生成与 AI 运维完成报告；统一导出合同", "reports page；inspection/task report APIs", "报告可追溯、可取消、可对账；导出一致"),
        ("T7 知识库", "检索、引用、来源质量、版本、删除和 AI 任务沉淀", "knowledge page；search/index APIs", "引用定位、隔离、删除与真实检索通过"),
        ("T8 设置", "接入、图谱、LLM、Agent/Workflow、MCP、RAG、策略、安全、平台健康", "settings sections；registry/graph/config/health APIs", "版本/探测/发布/回滚/审计；无 secret 到浏览器"),
        ("T9 能力覆盖", "把全部公共后端能力映射到八页或上下文动作；清理孤儿路由", "coverage ledger、router、API clients、contract tests", "零孤儿公共能力；零问题/资源/调查/处置独立入口"),
        ("T10 硬化", "可访问性、性能、安全、浏览器兼容、前端可观测性", "shared components、CSP、telemetry、tests", "AA audit、性能预算、安全回归"),
        ("T11 验收", "全页面视觉、功能、数据、后端覆盖、AI、灾难/恢复", "acceptance artifacts", "所有 P0/P1 门禁通过；零 mock；证据包可复核"),
    ], widths=[0.75, 2.05, 2.25, 1.55], font_size=7.8)
    add_heading(doc, "15.3 每个页面的实现模板", 2)
    add_numbered(doc, [
        "定义页面职责、入口、URL 状态、主动作和禁止项。",
        "定义每个字段的 API、源字段、聚合、单位、新鲜度和未知行为。",
        "先写单元/契约/路由/状态测试，再实现 view model 和组件。",
        "实现 loading、success、true empty、filter empty、partial、stale、unknown、not connected、failed 和 forbidden 全状态。",
        "接入统一 scope、request cancellation、错误诊断、审计和埋点。",
        "以真实环境完成 UI/API/source 三方核对并保存证据。",
        "按 UI 图做 1440×900、1024×768 和移动查看宽度视觉回归；修复后再进入下一页。",
        "把页面、字段、交互、后端能力、测试和证据写入能力覆盖账本；无账本记录不得声明完成。",
    ])
    add_heading(doc, "15.4 代码级强制检查", 2)
    add_bullets(doc, [
        "搜索并删除/替换生产路径中的 mock、fixture、Math.random、hard-coded health/count 和演示 fallback；测试目录可保留且必须隔离。",
        "搜索 `adminOnly`、旧路由 `view=`、hardware/vm/capacity/grafana/rules/telemetry 等重定向，逐一对照附录 A。",
        "搜索已定义但无调用的 API client 方法，逐一映射页面或明确弃用；构建时对 orphan capability 失败。",
        "对所有写 API 增加后端 scope 验证、幂等键、审计、预检和动作状态；前端确认只是体验层。",
        "对日志、AI prompt、知识索引、MCP 参数和错误消息实施秘密/PII 脱敏。",
    ])

    add_heading(doc, "16  完整验收门禁和交付物", 1)
    add_heading(doc, "16.1 验收环境纪律", 2)
    add_table(doc, ["要求", "判定"], [
        ("环境固定", "记录 commit/build、部署清单、依赖版本、浏览器、时区、租户、集群和测试窗口。"),
        ("只读优先", "发现/盘点阶段只读；任何状态变更先通过动作预检与确认。"),
        ("测试数据隔离", "测试对象带唯一 run ID 标签/前缀和 owner；只清理能证明归属的对象。"),
        ("真实源", "关闭 mock/fixture/demo fallback；网络请求和后端日志证明命中真实服务。"),
        ("可复现", "每个失败有 URL、步骤、scope、request/run ID、截图/视频、API/source 证据。"),
    ], widths=[1.35, 5.25])
    add_heading(doc, "16.2 六类发布门禁", 2)
    add_table(doc, ["门禁", "通过条件"], [
        ("G1 视觉", "目标分辨率无溢出/遮挡/错位；层级、间距、状态和交互一致；所有图例/单位可读。"),
        ("G2 功能", "所有入口、按钮、表单、过滤、深链、刷新、取消、导出、回退和异常恢复符合预期。"),
        ("G3 数据", "所有纳管集群和八页关键数值完成 UI/API/source 三方一致性核对；未知/部分/陈旧/失败语义正确；零生产 mock。"),
        ("G4 后端覆盖", "全接口形态的公共能力均归入 A/B/C 且真实可达；D 类仅内部使用；零孤儿入口和零孤儿公共能力。"),
        ("G5 AI", "能力图闭环可验证；40 个真实故障场景具备 ground truth、独立裁决和重复运行，全部达到质量与安全阈值。"),
        ("G6 非功能", "安全、租户隔离、危险动作、可访问性、响应式、性能、兼容、升级恢复、保留删除和平台自身可观测性通过。"),
    ], widths=[1.15, 5.45])
    add_heading(doc, "16.3 必测状态矩阵", 2)
    add_table(doc, ["对象", "状态"], [
        ("所有数据组件", "loading、success、true empty、filter empty、partial、stale、unknown、not connected、failed、forbidden"),
        ("账户与会话", "登录、退出、过期、撤销、刷新、跨标签页、直接深链、无角色切换器、完整导航一致"),
        ("范围与隔离", "租户/集群/资源切换、参数篡改、陈旧深链、并发切换、服务端拒绝、不可枚举其他租户"),
        ("列表", "0/1/少量/分页边界/大数据/重复/高基数字段/排序稳定性"),
        ("网络", "慢、超时、取消、断开、5xx、429、重试、乱序响应、部分成功"),
        ("实时/轮询", "首次加载、增量更新、重复事件、断线恢复、页面后台/恢复、scope 切换"),
        ("全链路监控", "无 Trace/日志/变更、部分采样、时钟偏差、服务图缺边、高基数服务、慢查询、异常区间切换"),
        ("写操作", "预检失败、对象漂移、确认取消、并发冲突、幂等重试、部分失败、回滚、验证失败"),
        ("图谱", "增量重复/乱序、全量漏项、checkpoint 恢复、tombstone、generation 发布失败、回滚和固定读取"),
        ("知识与报告", "索引陈旧/删除、引用失效、生成取消/失败、导出过期、页面与文件口径不一致"),
        ("AI", "无证据、冲突证据、工具超时、图谱 stale、知识缺失、提示注入、跨 scope 诱导、拒答"),
    ], widths=[1.45, 5.15])
    add_heading(doc, "16.4 自动化测试分层", 2)
    add_table(doc, ["层", "目标", "最低要求"], [
        ("静态", "类型、lint、路由和孤儿能力", "CI 必跑；生产 mock 扫描；公开 API 覆盖清单校验"),
        ("单元", "view model、状态机、格式化、权限/风险呈现", "关键分支和边界值"),
        ("契约", "前后端 schema、错误、scope、新鲜度", "provider/consumer contract + 真实样本"),
        ("组件", "加载/空/部分/失败、键盘、ARIA", "Story/fixture 只用于测试，不作为验收"),
        ("端到端", "跨页面主流程和写操作", "真实后端；核心 happy path + failure path"),
        ("数据对账", "UI/API/source 一致", "每个关键域至少一条自动或半自动核对"),
        ("视觉", "布局、断点、主题和长文本", "基准分辨率与关键状态截图 diff"),
        ("AI 评测", "正确性、证据、安全、稳定、时延", "40 场景 + 版本固定 + 结果可重放"),
        ("韧性/安全", "依赖失败、注入、跨租户、越权和灾难恢复", "零关键/高风险未解决"),
    ], widths=[1.0, 2.2, 3.4])
    doc.add_page_break()
    add_heading(doc, "16.5 缺陷严重度和处置", 2)
    add_body(doc, "设计优先级 P0/P1/P2 表示能力是否进入发布范围；缺陷严重度 S0/S1/S2/S3 表示一次失败的影响。两者不得混用。")
    add_table(doc, ["严重度", "判定", "发布处理"], [
        ("S0 阻断", "跨租户、越权写入、数据破坏、secret 泄露、AI 自动执行、不可恢复或验收证据造假", "立即停止相关测试和发布；修复并完成独立复验，不接受例外"),
        ("S1 严重", "八个核心页面不可用、关键数据错误、公共能力无归属、AI 定位/动作/验证主流程失败、无依据 Confirmed", "阻断发布；必须关闭并回归全部受影响流程"),
        ("S2 一般", "非核心功能失败、部分兼容/可访问性/性能不达标，存在明确安全替代路径", "仅可由正式例外接受；记录风险、范围、owner、期限和监控"),
        ("S3 轻微", "不影响理解和操作的视觉或文案偏差", "可进入后续修复，但不得累计破坏一致性或可访问性"),
    ], widths=[0.95, 3.85, 1.8], font_size=8.0)
    # Keep the evidence package together and away from the running footer.
    doc.add_page_break()
    add_heading(doc, "16.6 最终证据包", 2)
    add_bullets(doc, [
        "环境清单与版本：frontend/backend/orchestrator/graph/data source/build ID。",
        "页面清单：规范路由、关键状态、桌面/窄屏截图和视觉差异结果。",
        "功能结果：测试 ID、步骤、预期、实际、证据链接、owner、缺陷状态。",
        "数据对账：UI/API/source 值、查询、时间、scope、误差与结论。",
        "后端覆盖：REST/GraphQL/SSE/WebSocket/异步任务/定时任务/导出/回调 → A/B/C/D → UI path → 测试 → 结果。",
        "AI 报告：40 场景、版本组合、每次 run、证据、结论等级、阈值统计和失败样本。",
        "安全/非功能：隔离、危险动作、秘密、可访问性、性能、兼容和恢复测试。",
        "测试数据清理清单：仅测试所有对象、删除证据、未删除项及原因；严禁笼统“已清空”。",
        "最终签署：所有 P0/P1 能力通过，所有 S0/S1 缺陷关闭，S2 例外包含风险、范围、期限和 owner。",
    ])
    add_heading(doc, "16.7 运行态基线记录", 2)
    add_body(doc, "每次实施或验收开始时，都必须从真实 API 和来源系统生成当次运行态快照，记录 acceptance_run_id、完整时间与时区、scope、来源版本、request/run ID、字段口径、值和质量。历史截图、设计示例和上一次验收结果不能作为当前事实。若目标环境不可达，只记录 Not connected/Failed、本次失败证据与最后一次成功的独立时间戳；不得沿用旧数值填充当前页面或据此宣布通过。")
    add_table(doc, ["字段", "最低记录", "禁止"], [
        ("观测身份", "acceptance_run_id、执行者、build/commit、环境 URL", "无法关联到具体运行的截图"),
        ("时间与范围", "started_at/ended_at、时区、tenant/cluster/resource/time window", "只有相对时间或默认全平台"),
        ("数据事实", "UI 值、API 值、source 值、查询、误差、quality", "复制设计图数字或历史缓存冒充当前值"),
        ("不可达", "失败阶段、错误、request ID、最后成功时间、重试结果", "把未连接显示成空、0、正常或未知原因的陈旧值"),
    ], widths=[1.05, 3.9, 1.65], font_size=8.2)

    add_heading(doc, "17  其他智能体的执行提示词", 1)
    add_body(doc, "以下提示词可直接交给后续实现/验收智能体。方括号内容必须由执行环境实际值替换；不得删减“统一账户、真实数据、后端覆盖、AI 证据、安全门禁”约束。")
    prompt_lines = [
        "你是 AIOps 平台的实施与验收智能体。工作区为 [REPO_ROOT]，目标环境为 [BASE_URL]。",
        "首先完整阅读《docs/AIOps平台最终设计与实施规范.docx》。该文档是唯一产品基线；不得复用旧页面的信息架构，也不得按值班人员/管理员拆分产品。",
        "在修改前：检查仓库说明和未提交改动；建立唯一 acceptance_run_id 与证据目录；读取前端路由、REST/GraphQL、SSE/WebSocket、异步/定时任务、导出、回调、真实数据源和当前环境；生成第 8.4 节规定的能力覆盖账本及差距。",
        "按 T0—T11 顺序实施。以图 1—8 为八个页面结构基线；每个页面先建立字段合同与测试，再实现 loading、success、true empty、filter empty、partial、stale、unknown、not connected、failed 和 forbidden。后端缺口必须写成契约并在后端实现，不得用 mock、fixture、硬编码或随机数据填充生产界面。逐页执行第 5.4 节信息有效性审计：删除无口径计数、重复健康卡片、无动作标签、历史值和不可打开证据。",
        "统一账户主导航严格按顺序提供 AI 智能运维、总览、集群、全链路监控、知识图谱、知识库、报告、设置，不得新增问题、资源、调查、处置入口。总览汇总全部纳管集群告警与集群级关键指标；集群页展示指定集群集群级 CPU/内存、热点节点、Kubernetes、网络和存储；全链路只使用云平台控制面/计算/网络/存储/Kubernetes 路径并统一绑定 selectedPathId；知识图谱统一绑定 selectedRootEntityId、selectedRelationId 和 generation。",
        "AI 智能运维是默认首页，必须在一页完成范围冻结、证据、假设、反证、结论等级、动作草稿、预检、显式确认、执行、回滚和恢复验证。告警、资源、调查阶段与处置记录只是八页中的上下文对象。危险操作需后端范围校验、影响清单、幂等、审计和恢复验证；AI 不能自动批准或持有基础设施凭据。",
        "报告页只设巡检报告和 AI 运维报告两类主入口。巡检覆盖云平台告警、计算、Kubernetes、网络、存储、容量、数据质量与平台自身健康；AI 报告从已完成/终止 task run 生成并绑定证据、动作和前后验证。页面与导出使用同一查询合同。",
        "知识图谱必须验证自动更新闭环：八类来源、cadence、阶段、quality、last attempted/last success、generation、计数、错误、对账、幂等重放、checkpoint、tombstone、原子发布、回滚，以及总览/集群/全链路/知识图谱/AI 对同一 generation 的消费。原始 JSON 仅是诊断附属视图。",
        "设置页必须覆盖接入、图谱运维、LLM 路由、Agent、Skill、Workflow、MCP server/tools、知识与 RAG、策略安全、评测和平台运行健康。每次生产 run 必须可复现模型、工作流、知识、图谱和策略版本。",
        "智能故障定位必须按 scope 冻结→证据→假设→反证/因果验证→结论等级→动作草稿→执行验证闭环工作。按第 11.5 节冻结 ground truth、隔离答案、双人独立评审和第三人裁决；每个场景运行 5 次。证据不足不得输出 Confirmed。运行 40 个真实故障场景并达到文档阈值。",
        "验收使用真实环境。每次先按第 16.7 节生成当次运行态基线；环境不可达时只记录失败与最后成功时间，不得沿用历史值。仅删除能通过 run ID/标签/前缀证明属于本次测试的数据，不得广泛删除。每个关键字段提交 UI/API/source 三方对账。",
        "每完成一个阶段运行对应单元、契约、端到端、视觉、数据和安全测试，保存可复核证据。发现失败时定位根因并修复，不得仅隐藏错误或放宽断言。",
        "完成前逐项执行 G1—G6 门禁并执行信息效率零容忍检查：八个一级页面均真实可达且顺序正确；没有问题/资源/调查/处置独立入口；所有纳管集群在总览有明确状态；全链路和图谱选择一致；零无口径 KPI、零硬编码可变运行值、零图谱计数矛盾、零不可打开 AI 证据、零跨范围标识泄露。只有全部 P0/P1 能力通过、S0/S1 缺陷关闭，且没有模拟数据、孤儿公共能力、无效入口、AI 自动确认或敏感信息泄露时，才可声明完成。",
    ]
    for i, line in enumerate(prompt_lines, 1):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.25)
        p.paragraph_format.first_line_indent = Inches(-0.25)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.0
        add_text(p, f"{i}. ", bold=True, color=BLUE)
        add_text(p, line, size=8.5)

    add_heading(doc, "附录 A  路由和旧入口迁移表", 1)
    add_table(doc, ["旧入口/模式", "目标", "参数转换/处理"], [
        ("/ 或旧 Assistant/Chat", "/ai-operations", "replace；保留 conversationId、scope、time 和可转换上下文"),
        ("/clusters/:id/resources", "/clusters/:clusterUid/overview?tab=resources", "将旧集群 ID 解析为 canonical cluster UID；保留资源筛选"),
        ("observe/metrics/traces/logs", "/observability", "转换 cluster/service/time/signal；无法转换时显示明确迁移说明"),
        ("problems/alerts", "/overview 或 /clusters/:clusterUid/overview", "按 scope 跳转；携带 alert/event ID；不得保留独立问题页"),
        ("resources", "/clusters/:clusterUid/overview 或 /knowledge-graph", "按 resourceUid、domain、关系意图选择目标；不得保留独立资源页"),
        ("investigate", "/ai-operations", "保留 task/run/investigation ID；映射为 AI 任务阶段"),
        ("actions/ops", "/ai-operations", "保留 actionId/status；映射为任务内动作生命周期"),
        ("knowledge", "/knowledge", "保留 source/query"),
        ("reports / dashboards", "/reports", "保留 reportId/type/scope/time；仅巡检/AI 运维两类主入口"),
        ("capacity → resources", "/clusters/:clusterUid/overview?section=capacity", "保留 resourceUid/domain；目标 section 必须真实存在"),
        ("grafana/telemetry", "/observability", "转换为 selectedPathId/scope/time；无 silent fallback"),
        ("alert rules / silences", "/settings?section=integrations", "映射到信号治理分区；写操作保留审计和 scope"),
        ("hardware/vm redirect", "/clusters/:clusterUid/overview?section=compute", "保留 resourceUid/type；可进一步打开图谱或 AI"),
        ("system/admin/capabilities", "/settings", "不按角色隐藏；按同页 section 参数定位"),
        ("cluster/source settings", "/settings?section=integrations", "统一接入、探测、目录和凭据引用"),
        ("graph raw ops", "/settings?section=graph", "同步源、质量、版本、对账与恢复；raw JSON 为诊断附属视图"),
        ("graph explorer", "/knowledge-graph", "携带 selectedRootEntityId、selectedRelationId、time、generation"),
        ("AI capability/config", "/settings?section=llm|workflow|mcp|rag", "按原对象 ID 和版本映射到对应设置分区"),
        ("security/scope/credentials", "/settings?section=security", "统一账户可见；服务端策略决定读取和写操作"),
        ("components/system health", "/settings?section=health", "按端到端能力链显示真实探测与用户影响"),
    ], widths=[2.0, 2.1, 2.5], font_size=7.8)

    add_heading(doc, "附录 B  页面级验收检查表", 1)
    checks = [
        ("全局壳层", "统一账户看到八个一级入口且顺序固定，AI 智能运维第一；无问题/资源/调查/处置导航；scope/time 可深链；无 adminOnly；顶栏标明数据截止、时区与质量。"),
        ("AI 智能运维", "默认首页可用；冻结上下文、证据定位、假设/反证、结论等级、动作草稿、预检、确认、执行、回滚、验证和报告生成均可复核。"),
        ("总览", "所有纳管集群均出现；告警、集群级 CPU/内存、节点/Pod 就绪、网络、存储、容量和数据质量有明确值或状态；图表可下钻。"),
        ("集群", "指定 clusterUid 生效；集群 CPU/内存口径、热点节点、Pod 调度/重启、网络和存储指标有单位/时间/来源/质量；活动告警进入 AI。"),
        ("全链路监控", "仅展示云平台路径；推荐原因明确；所有面板绑定同一 selectedPathId；依赖、风险排行、趋势、性能矩阵和证据可打开。"),
        ("知识图谱", "中心对象和选中关系明确；所有面板绑定 selectedRootEntityId/selectedRelationId/generation；两跳上限、影响分析、关系事实和自动更新状态可验证。"),
        ("知识库", "来源、索引、版本、新鲜度、原文引用、任务沉淀审阅、删除和租户隔离可验证。"),
        ("报告", "巡检报告和 AI 运维报告均来自真实数据；计划、生成、取消、失败重试、版本追溯及页面/导出对账通过。"),
        ("设置", "接入、图谱、LLM、Agent/Skill、Workflow、MCP、RAG、策略、安全、评测和平台健康均有真实状态、版本、探测、发布/回滚和审计。"),
        ("后端覆盖", "每个公共能力归入八页或上下文动作；零孤儿路由、零仅 API client 覆盖声明；内部能力在设置健康分区可观测。"),
        ("信息有效性", "零无口径 KPI/徽标；健康项异常优先折叠；完整时间与时区；图谱总数/渲染数一致；AI 证据均可打开；不可达不复用历史值。"),
        ("非功能", "AA、响应式、性能预算、兼容、前端错误关联、秘密/跨范围/注入/危险动作测试通过。"),
    ]
    add_table(doc, ["页面/域", "必须通过"], checks, widths=[1.35, 5.25])

    add_heading(doc, "附录 C  需求追踪矩阵和最终完成定义", 1)
    add_table(doc, ["需求", "设计落点", "验收证据"], [
        ("页面美观正常", "第 5、6、14 章与图 1—8", "视觉回归、分辨率矩阵、AA 与性能报告"),
        ("AI 第一入口", "图 1、第 4、6.1、11 章", "默认路由、冻结上下文、证据、结论、动作与验证 E2E"),
        ("全局集群态势", "图 2、第 4、6.2 章", "纳管集群清单、告警和集群级关键指标 UI/API/source 对账"),
        ("指定集群总览", "图 3、第 6.3、7 章", "clusterUid 深链；集群级 CPU/内存、热点、Kubernetes/网络/存储及告警对账"),
        ("全链路监控", "图 4、第 6.4、7.3 章", "云平台路径、selectedPathId 一致性、性能矩阵和可打开证据"),
        ("知识图谱独立页面", "图 5、第 6.5、9 章", "中心/关系选择一致、有限子图、影响路径、generation 与业务下钻"),
        ("知识库独立页面", "图 6、第 6.6 章", "来源、索引、版本、新鲜度、引用和沉淀审阅"),
        ("两类报告", "图 7、第 6.7 章", "巡检与 AI 运维报告生成、计划、追溯和导出对账"),
        ("统一设置", "图 8、第 6.8、9—10 章", "LLM/Workflow/MCP/RAG/图谱/接入/安全/健康的真实状态与版本"),
        ("所有前端功能正常", "第 4、6、15、16 章", "路由/组件/E2E/失败恢复结果"),
        ("真实环境数据", "第 7、12、16 章", "零 mock 扫描；UI/API/source 三方对账"),
        ("前端覆盖后端", "第 8 章与附录 A", "全接口形态能力覆盖账本；零孤儿公共能力"),
        ("AI 智能故障定位", "图 1、图 10、第 10—11 章", "40 场景、ground truth、重复运行、裁决、阈值和安全测试"),
        ("知识图谱自动更新/展示", "图 5、图 9 与第 9 章", "八源调度、quality/generation、重放/回滚、对账和业务消费"),
        ("云平台运维信息全面", "第 3、7 章", "五域资产、八类信号、可靠性与完整性状态"),
        ("智能能力全面", "图 8、第 10 章", "LLM/Workflow/MCP/Agent/Skill/RAG/策略/评测运行"),
        ("统一账户", "第 2、4、13 章", "完整导航一致；无角色拆页；写操作仍受控"),
        ("信息有效且高效", "第 5.4、6、16.7 章", "零无口径/重复/历史伪当前信息；图谱计数一致；证据可打开；时间与范围明确"),
    ], widths=[1.55, 2.55, 2.5], font_size=8.5)
    add_heading(doc, "C.1 最终完成定义", 2)
    add_body(doc, "只有同时满足以下条件才可宣布“完整验收通过”：")
    add_bullets(doc, [
        "图 1—8 的信息层级、主流程、状态语义和统一账户约束实现完成；八项一级导航顺序固定且 AI 智能运维为默认首页；无问题、资源、调查、处置独立入口。",
        "REST、GraphQL、SSE、WebSocket、异步任务、定时任务、导出和回调等公共能力可从 A/B/C 入口真实到达并验证；内部 D 类有平台健康与运行诊断。",
        "所有关键页面使用真实环境数据，且完成 UI/API/source 三方对账；生产路径无模拟数据。",
        "逐页通过信息有效性检查：零无口径徽标/KPI、零硬编码可变运行值、零跨范围标识泄露、零图谱计数矛盾、零不可打开 AI 证据；不可达状态不复用历史值。",
        "图谱八类来源自动更新、质量发布、generation、对账和业务消费闭环可复核。",
        "设置中的 AI 能力配置完整，智能故障定位达到第 11 章阈值，零跨范围证据、零无依据 Confirmed、零 AI 自动确认或越权写入。",
        "危险操作、秘密、租户隔离、提示注入、可访问性、性能和恢复测试全部通过。",
        "所有 P0/P1 能力通过，所有 S0/S1 缺陷关闭；证据包可由另一智能体在同环境独立复现。",
    ])

    # Closing revision record.
    add_heading(doc, "修订记录", 1)
    add_table(doc, ["版本", "日期", "变更"], [
        ("V1.4", "2026-09-13", "固化八页最终信息架构：AI 智能运维作为第一入口；保留总览、集群、全链路、知识图谱、知识库、报告、设置；取消问题/资源/调查/处置独立页面；补齐云平台路径选择、图谱选择、两类报告及统一设置规范，并替换全部页面设计图。"),
        ("V1.3", "2026-09-12", "按五页观察架构复审：总览汇总全部纳管集群告警和关键指标；新增指定集群总览与全链路监控；知识图谱和知识库升级为独立一级入口；补齐路由、数据合同、后端归属、实施任务和验收矩阵。"),
        ("V1.2", "2026-09-12", "再次自审：删除无口径导航计数与低效健康卡片；明确数据截止/时区；修正 NAT；补齐图谱渲染计数、AI 证据定位、租户不可枚举、不可达状态和信息效率门禁。"),
        ("V1.1", "2026-09-12", "自审核修订：分离知识与报告入口；新增七个页面级 UI；补齐全接口覆盖账本、图谱恢复、AI ground truth 评测、兼容生命周期和缺陷严重度。"),
        ("V1.0", "2026-09-12", "从零建立最终产品、UI、数据、后端覆盖、图谱、AI、安全、实施与验收规范；固化统一账户。"),
    ], widths=[0.9, 1.2, 4.5])

    # Prevent the cover's H1 page-break behavior from making the first navigation awkward is intentional.
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build_doc()
